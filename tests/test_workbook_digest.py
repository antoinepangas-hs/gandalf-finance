"""Tests for the deterministic workbook digest used in the judge prompt."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.workbook.defined_name import DefinedName  # type: ignore[import-untyped]

from gandalf.workbook_digest import build_workbook_digest


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Revenue"
    summary["B1"] = "Cost"
    summary["A2"] = 100
    summary["B2"] = "=A2*1.1"
    summary["B3"] = "=SUM(B1:B2)"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Secret"
    workbook.defined_names["Revenue_Cell"] = DefinedName("Revenue_Cell", attr_text="Summary!$A$2")
    workbook.save(path)


def test_digest_reports_sheets_headers_formulas_and_names(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    _make_workbook(workbook)

    digest = build_workbook_digest(workbook)

    assert digest is not None
    assert "Sheets (2):" in digest
    assert '"Summary" (visible)' in digest
    assert '"Hidden" (hidden)' in digest
    assert "row1: Revenue | Cost" in digest
    assert "formula B2: =A2*1.1" in digest
    assert "formula B3: =SUM(B1:B2)" in digest
    assert "Revenue_Cell -> Summary!$A$2" in digest


def test_digest_respects_formula_sample_cap(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 11):
        ws.cell(row=row, column=1, value=f"=ROW()+{row}")
    wb.save(workbook)

    digest = build_workbook_digest(workbook, max_formula_samples=3, max_preview_columns=0)

    assert digest is not None
    assert digest.count("formula ") == 3


def test_digest_truncates_sheet_list(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    wb = Workbook()
    wb.active.title = "S0"
    for i in range(1, 5):
        wb.create_sheet(f"S{i}")
    wb.save(workbook)

    digest = build_workbook_digest(workbook, max_sheets=2)

    assert digest is not None
    assert "and 3 more sheet(s) not shown" in digest


def test_digest_returns_none_for_unsupported_extension(tmp_path: Path) -> None:
    legacy = tmp_path / "model.xls"
    legacy.write_bytes(b"not a real xls")
    assert build_workbook_digest(legacy) is None


def test_digest_returns_none_for_unreadable_file(tmp_path: Path) -> None:
    broken = tmp_path / "model.xlsx"
    broken.write_bytes(b"not a real xlsx")
    assert build_workbook_digest(broken) is None


def test_digest_returns_none_for_missing_file(tmp_path: Path) -> None:
    assert build_workbook_digest(tmp_path / "nope.xlsx") is None
