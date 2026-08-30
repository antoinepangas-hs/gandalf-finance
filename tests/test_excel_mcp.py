"""Tests for the Excel-backed MCP service layer."""

from __future__ import annotations

import base64
import hashlib
import json
import os
import shutil
import subprocess
import threading
import time
import urllib.error
from collections.abc import Iterator
from contextlib import contextmanager
from email.message import Message
from http.server import ThreadingHTTPServer
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from gandalf import excel_mcp
from gandalf.excel_mcp import (
    ExcelServerConfig,
    ExcelService,
    LibreOfficeBackend,
    WindowsVMExcelBackend,
    build_worker_server,
    cell_address,
    config_from_env,
    handle_worker_ensure,
    handle_worker_request,
    prune_worker_cache,
    range_cell_count,
    range_number_formats,
    run_worker,
    worker_cache_lock,
    worker_main,
    worker_request_is_authorised,
)


class FakeBackend:
    """Fake backend used to test validation and stable service responses."""

    def workbook_summary(self, workbook_path: Path, *, recalculate: bool) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "fake",
            "recalculated": recalculate,
            "sheets": [{"name": "Sheet1", "used_range": "$A$1:$B$2"}],
        }

    def inspect_range(
        self,
        workbook_path: Path,
        *,
        sheet: str,
        range_address: str,
        include_values: bool,
        include_formulas: bool,
        include_formats: bool,
        recalculate: bool,
    ) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "fake",
            "sheet": sheet,
            "range": range_address,
            "values": [[1, 2]] if include_values else None,
            "formulas": [["=1", "=2"]] if include_formulas else None,
            "number_formats": [["0", "0"]] if include_formats else None,
            "recalculated": recalculate,
        }

    def find_cells(
        self,
        workbook_path: Path,
        *,
        query: str,
        search_values: bool,
        search_formulas: bool,
        max_results: int,
        max_cells_per_call: int,
    ) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "fake",
            "query": query,
            "matches": [
                {
                    "sheet": "Sheet1",
                    "cell": "A1",
                    "match_type": "formula" if search_formulas else "value",
                    "value": "Revenue",
                    "formula": "=SUM(B1:B2)" if search_formulas else None,
                }
            ][:max_results],
            "searched_values": search_values,
            "max_cells_per_call": max_cells_per_call,
        }

    def recalculate(self, workbook_path: Path) -> dict[str, Any]:
        return {"path": workbook_path.name, "backend": "fake", "recalculated": True}

    def workbook_repair_check(self, workbook_path: Path) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "fake",
            "checked": True,
            "opened": True,
            "repair_dialog_detected": False,
        }


def make_service(tmp_path: Path, *, allow_macros: bool = False, max_cells_per_call: int = 10) -> ExcelService:
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (workdir / "model.xlsx").write_bytes(b"not really an xlsx")
    return ExcelService(
        workdir=workdir,
        backend=FakeBackend(),
        allow_macros=allow_macros,
        max_cells_per_call=max_cells_per_call,
    )


def make_openpyxl_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet["A1"] = "Revenue"
    worksheet["B1"] = "=SUM(B2:B3)"
    worksheet["B2"] = 600
    worksheet["B3"] = 400
    worksheet["C1"] = 0.123
    worksheet["C1"].number_format = "0.0%"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Private"
    workbook.save(path)


def test_service_rejects_path_traversal(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(b"outside")

    with pytest.raises(ValueError, match="outside the allowed workdir"):
        service.workbook_summary(str(outside))


def test_service_rejects_macro_enabled_workbooks_by_default(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    macro_file = tmp_path / "workdir" / "model.xlsm"
    macro_file.write_bytes(b"macro workbook")

    with pytest.raises(ValueError, match="may contain macros"):
        service.workbook_summary("model.xlsm")


def test_service_allows_macro_enabled_workbooks_when_configured(tmp_path: Path) -> None:
    service = make_service(tmp_path, allow_macros=True)
    macro_file = tmp_path / "workdir" / "model.xlsm"
    macro_file.write_bytes(b"macro workbook")

    result = service.workbook_summary("model.xlsm")

    assert result["path"] == "model.xlsm"


def test_service_enforces_range_cell_limit(tmp_path: Path) -> None:
    service = make_service(tmp_path, max_cells_per_call=2)

    with pytest.raises(ValueError, match="exceeding max_cells_per_call"):
        service.inspect_range("model.xlsx", sheet="Sheet1", range_address="A1:B2")


def test_service_enforces_smaller_format_cell_limit(tmp_path: Path) -> None:
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (workdir / "model.xlsx").write_bytes(b"not really an xlsx")
    service = ExcelService(
        workdir=workdir,
        backend=FakeBackend(),
        allow_macros=False,
        max_cells_per_call=10,
        max_format_cells_per_call=1,
    )

    result = service.inspect_range(
        "model.xlsx",
        sheet="Sheet1",
        range_address="A1:B1",
        include_formats=False,
    )
    assert result["values"] == [[1, 2]]

    with pytest.raises(ValueError, match="max_format_cells_per_call"):
        service.inspect_range("model.xlsx", sheet="Sheet1", range_address="A1:B1", include_formats=True)


def test_service_returns_stable_summary_range_find_and_recalculate(tmp_path: Path) -> None:
    service = make_service(tmp_path)

    assert service.workbook_summary("model.xlsx")["sheets"] == [{"name": "Sheet1", "used_range": "$A$1:$B$2"}]
    range_result = service.inspect_range("model.xlsx", sheet="Sheet1", range_address="A1:B1")
    assert range_result["values"] == [[1, 2]]
    assert range_result["formulas"] == [["=1", "=2"]]
    assert range_result["number_formats"] is None
    find_result = service.find_cells("model.xlsx", query="Revenue")
    assert find_result["matches"][0]["cell"] == "A1"
    assert service.recalculate("model.xlsx")["recalculated"] is True
    assert service.workbook_repair_check("model.xlsx")["repair_dialog_detected"] is False


def test_service_caches_summary_range_and_find_results(tmp_path: Path) -> None:
    class CountingBackend(FakeBackend):
        def __init__(self) -> None:
            self.summary_calls = 0
            self.range_calls = 0
            self.find_calls = 0

        def workbook_summary(self, workbook_path: Path, *, recalculate: bool) -> dict[str, Any]:
            self.summary_calls += 1
            return super().workbook_summary(workbook_path, recalculate=recalculate)

        def inspect_range(self, workbook_path: Path, **kwargs: Any) -> dict[str, Any]:
            self.range_calls += 1
            return super().inspect_range(workbook_path, **kwargs)

        def find_cells(self, workbook_path: Path, **kwargs: Any) -> dict[str, Any]:
            self.find_calls += 1
            return super().find_cells(workbook_path, **kwargs)

    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (workdir / "model.xlsx").write_bytes(b"not really an xlsx")
    metrics_path = tmp_path / "excel_metrics.json"
    backend = CountingBackend()
    service = ExcelService(
        workdir=workdir,
        backend=backend,
        allow_macros=False,
        max_cells_per_call=10,
        metrics_path=metrics_path,
    )

    first_summary = service.workbook_summary("model.xlsx")
    second_summary = service.workbook_summary("model.xlsx")
    first_range = service.inspect_range("model.xlsx", sheet="Sheet1", range_address="A1:B1")
    second_range = service.inspect_range("model.xlsx", sheet="Sheet1", range_address="A1:B1")
    first_find = service.find_cells("model.xlsx", query="Revenue")
    second_find = service.find_cells("model.xlsx", query="Revenue")

    assert backend.summary_calls == 1
    assert backend.range_calls == 1
    assert backend.find_calls == 1
    assert "metrics" not in first_summary
    assert "metrics" not in second_summary
    assert "metrics" not in first_range
    assert "metrics" not in second_range
    assert "metrics" not in first_find
    assert "metrics" not in second_find
    metrics = json.loads(metrics_path.read_text())
    assert metrics["call_count"] == 6
    assert metrics["operation_counts"] == {
        "find_cells": 2,
        "inspect_range": 2,
        "workbook_summary": 2,
    }
    assert metrics["excel_service_cache_hits"] == 3
    assert metrics["excel_service_cache_misses"] == 3

    service.recalculate("model.xlsx")
    service.workbook_summary("model.xlsx")

    assert backend.summary_calls == 2
    metrics = json.loads(metrics_path.read_text())
    assert metrics["call_count"] == 8
    assert metrics["operation_counts"]["recalculate"] == 1
    assert metrics["operation_counts"]["workbook_summary"] == 3


def test_service_debug_metrics_keeps_metrics_in_response(tmp_path: Path) -> None:
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (workdir / "model.xlsx").write_bytes(b"not really an xlsx")
    service = ExcelService(
        workdir=workdir,
        backend=FakeBackend(),
        allow_macros=False,
        max_cells_per_call=10,
        debug_metrics=True,
    )

    result = service.workbook_summary("model.xlsx")

    assert result["metrics"]["excel_service_cache_hit"] is False


def test_service_strips_metrics_by_default_and_aggregates_to_file(tmp_path: Path) -> None:
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (workdir / "model.xlsx").write_bytes(b"not really an xlsx")
    metrics_path = tmp_path / "excel_metrics.json"
    service = ExcelService(
        workdir=workdir,
        backend=FakeBackend(),
        allow_macros=False,
        max_cells_per_call=10,
        metrics_path=metrics_path,
    )

    first = service.workbook_summary("model.xlsx")
    second = service.workbook_summary("model.xlsx")  # served from the operation cache

    assert "metrics" not in first  # not exposed to the agent by default
    assert "metrics" not in second
    summary = json.loads(metrics_path.read_text())
    assert summary["call_count"] == 2
    assert summary["operation_counts"]["workbook_summary"] == 2
    assert summary["excel_service_cache_misses"] == 1
    assert summary["excel_service_cache_hits"] == 1


def test_merge_excel_metrics_summaries_totals_counts_and_sums() -> None:
    first = {
        "call_count": 2,
        "operation_counts": {"workbook_summary": 2},
        "excel_service_cache_hits": 1,
        "windows_vm_client_payload_bytes": 100.0,
    }
    second = {
        "call_count": 3,
        "operation_counts": {"workbook_summary": 1, "find_cells": 2},
        "excel_service_cache_hits": 2,
        "windows_vm_client_payload_bytes": 50.0,
    }

    merged = excel_mcp.merge_excel_metrics_summaries([first, second, "not a dict"])  # type: ignore[list-item]

    assert merged["call_count"] == 5
    assert merged["operation_counts"] == {"workbook_summary": 3, "find_cells": 2}
    assert merged["excel_service_cache_hits"] == 3
    assert merged["windows_vm_client_payload_bytes"] == 150.0


def test_libreoffice_backend_summary_and_range_use_openpyxl_without_recalculate(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    backend = LibreOfficeBackend(soffice_path="/does/not/exist")

    summary = backend.workbook_summary(workbook, recalculate=False)
    inspected = backend.inspect_range(
        workbook,
        sheet="Summary",
        range_address="A1:C1",
        include_values=True,
        include_formulas=True,
        include_formats=True,
        recalculate=False,
    )

    assert summary["backend"] == "libreoffice"
    assert summary["converted"] is False
    assert summary["recalculated"] is False
    assert summary["sheets"] == [
        {"name": "Summary", "used_range": "A1:C3", "rows": 3, "columns": 3, "visible": True},
        {"name": "Hidden", "used_range": "A1:A1", "rows": 1, "columns": 1, "visible": False},
    ]
    assert inspected["converted"] is False
    assert inspected["recalculated"] is False
    assert inspected["values"][0][0] == "Revenue"
    assert inspected["formulas"] == [["Revenue", "=SUM(B2:B3)", 0.123]]
    assert inspected["number_formats"] == [["General", "General", "0.0%"]]


def test_libreoffice_backend_reuses_loaded_openpyxl_workbooks(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    original_load = excel_mcp.load_openpyxl_workbook
    load_calls: list[bool] = []

    def tracking_load(path: Path, *, data_only: bool) -> Any:
        load_calls.append(data_only)
        return original_load(path, data_only=data_only)

    monkeypatch.setattr(excel_mcp, "load_openpyxl_workbook", tracking_load)
    backend = LibreOfficeBackend(soffice_path="/does/not/exist")

    first = backend.inspect_range(
        workbook,
        sheet="Summary",
        range_address="A1:B1",
        include_values=True,
        include_formulas=True,
        include_formats=False,
        recalculate=False,
    )
    second = backend.inspect_range(
        workbook,
        sheet="Summary",
        range_address="B2:B3",
        include_values=True,
        include_formulas=True,
        include_formats=False,
        recalculate=False,
    )

    assert load_calls == [False, True]
    assert first["metrics"]["libreoffice_loaded_workbook_cache_hit"] is False
    assert second["metrics"]["libreoffice_loaded_workbook_cache_hit"] is True


def test_libreoffice_cache_metrics_are_aggregated_by_service(tmp_path: Path) -> None:
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    workbook = workdir / "model.xlsx"
    make_openpyxl_workbook(workbook)
    metrics_path = tmp_path / "excel_metrics.json"
    service = ExcelService(
        workdir=workdir,
        backend=LibreOfficeBackend(soffice_path="/does/not/exist"),
        allow_macros=False,
        max_cells_per_call=20,
        metrics_path=metrics_path,
    )

    service.inspect_range("model.xlsx", sheet="Summary", range_address="A1:B1", recalculate=False)
    service.inspect_range("model.xlsx", sheet="Summary", range_address="B2:B3", recalculate=False)

    metrics = json.loads(metrics_path.read_text())
    assert metrics["libreoffice_loaded_workbook_cache_misses"] == 1
    assert metrics["libreoffice_loaded_workbook_cache_hits"] == 1
    assert metrics["libreoffice_openpyxl_load_ms"] >= 0


def test_libreoffice_backend_find_and_recalculate_use_converted_copy(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    calls: list[Path] = []

    def fake_convert(
        _backend: LibreOfficeBackend,
        workbook_path: Path,
        temp_dir: Path,
        *,
        force_recalculate: bool,
    ) -> Path:
        calls.append(workbook_path)
        assert force_recalculate is True
        converted = temp_dir / workbook_path.name
        shutil.copy2(workbook_path, converted)
        return converted

    monkeypatch.setattr(LibreOfficeBackend, "_convert_with_libreoffice", fake_convert)
    backend = LibreOfficeBackend(soffice_path="/does/not/exist")

    find_result = backend.find_cells(
        workbook,
        query="SUM",
        search_values=False,
        search_formulas=True,
        max_results=10,
        max_cells_per_call=20,
    )
    recalc_result = backend.recalculate(workbook)

    assert find_result["backend"] == "libreoffice"
    assert find_result["converted"] is True
    assert find_result["recalculated"] is True
    assert find_result["matches"] == [
        {
            "sheet": "Summary",
            "cell": "B1",
            "match_type": "formula",
            "value": None,
            "formula": "=SUM(B2:B3)",
        }
    ]
    assert recalc_result["path"] == "model.xlsx"
    assert recalc_result["backend"] == "libreoffice"
    assert recalc_result["converted"] is True
    assert recalc_result["recalculated"] is True
    assert recalc_result["metrics"]["libreoffice_prepared_cache_hit"] is True
    assert calls == [workbook]


def test_libreoffice_backend_reuses_prepared_file_across_instances_and_clone_paths(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    clone = tmp_path / "clone" / "model.xlsx"
    clone.parent.mkdir()
    make_openpyxl_workbook(workbook)
    shutil.copy2(workbook, clone)
    cache_dir = tmp_path / "lo-cache"
    calls: list[Path] = []

    def fake_convert(
        _backend: LibreOfficeBackend,
        workbook_path: Path,
        temp_dir: Path,
        *,
        force_recalculate: bool,
    ) -> Path:
        calls.append(workbook_path)
        assert force_recalculate is True
        converted = temp_dir / workbook_path.name
        shutil.copy2(workbook_path, converted)
        return converted

    monkeypatch.setattr(LibreOfficeBackend, "_convert_with_libreoffice", fake_convert)

    first = LibreOfficeBackend(soffice_path="/does/not/exist", cache_dir=cache_dir).recalculate(workbook)
    second = LibreOfficeBackend(soffice_path="/does/not/exist", cache_dir=cache_dir).recalculate(clone)

    assert calls == [workbook]
    assert first["metrics"]["libreoffice_prepared_cache_hit"] is False
    assert second["metrics"]["libreoffice_prepared_cache_hit"] is True


def test_libreoffice_backend_changed_workbook_bytes_miss_prepared_cache(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    cache_dir = tmp_path / "lo-cache"
    calls = 0

    def fake_convert(
        _backend: LibreOfficeBackend,
        workbook_path: Path,
        temp_dir: Path,
        *,
        force_recalculate: bool,
    ) -> Path:
        nonlocal calls
        calls += 1
        converted = temp_dir / workbook_path.name
        shutil.copy2(workbook_path, converted)
        return converted

    monkeypatch.setattr(LibreOfficeBackend, "_convert_with_libreoffice", fake_convert)

    backend = LibreOfficeBackend(soffice_path="/does/not/exist", cache_dir=cache_dir)
    first = backend.recalculate(workbook)
    workbook.write_bytes(workbook.read_bytes() + b"changed")
    second = backend.recalculate(workbook)

    assert calls == 2
    assert first["metrics"]["libreoffice_prepared_cache_hit"] is False
    assert second["metrics"]["libreoffice_prepared_cache_hit"] is False


def test_libreoffice_backend_publishes_prepared_file_atomically(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    cache_dir = tmp_path / "lo-cache"
    original_replace = os.replace
    replace_calls: list[tuple[Path, Path]] = []

    def fake_convert(
        _backend: LibreOfficeBackend,
        workbook_path: Path,
        temp_dir: Path,
        *,
        force_recalculate: bool,
    ) -> Path:
        converted = temp_dir / workbook_path.name
        converted.write_bytes(b"complete prepared workbook")
        return converted

    def tracking_replace(source: Any, dest: Any) -> None:
        source_path = Path(source)
        dest_path = Path(dest)
        assert dest_path.name == "prepared.xlsx"
        assert not dest_path.exists()
        original_replace(source_path, dest_path)
        replace_calls.append((source_path, dest_path))

    monkeypatch.setattr(LibreOfficeBackend, "_convert_with_libreoffice", fake_convert)
    monkeypatch.setattr("gandalf.excel_mcp.os.replace", tracking_replace)

    result = LibreOfficeBackend(soffice_path="/does/not/exist", cache_dir=cache_dir).recalculate(workbook)

    prepared_path = replace_calls[0][1]
    assert result["metrics"]["libreoffice_prepared_cache_hit"] is False
    assert len(replace_calls) == 1
    assert prepared_path.read_bytes() == b"complete prepared workbook"


def test_libreoffice_backend_memoizes_hash_for_unchanged_source(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    calls = 0
    original_hash = excel_mcp.file_sha256

    def tracking_hash(path: Path) -> str:
        nonlocal calls
        calls += 1
        return original_hash(path)

    monkeypatch.setattr(excel_mcp, "file_sha256", tracking_hash)
    backend = LibreOfficeBackend(soffice_path="/does/not/exist")

    first_key, _first_metrics = backend._cache_key(workbook, force_recalculate=True)
    second_key, _second_metrics = backend._cache_key(workbook, force_recalculate=True)

    assert first_key == second_key
    assert calls == 1


def test_libreoffice_backend_loaded_workbook_lru_eviction_closes_workbooks(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    first.write_bytes(b"first")
    second.write_bytes(b"second")
    closed: list[str] = []

    class FakeWorkbook:
        def __init__(self, name: str) -> None:
            self.name = name

        def close(self) -> None:
            closed.append(self.name)

    def fake_load(path: Path, *, data_only: bool) -> FakeWorkbook:
        return FakeWorkbook(f"{path.name}:{data_only}")

    monkeypatch.setattr(excel_mcp, "load_openpyxl_workbook", fake_load)
    backend = LibreOfficeBackend(soffice_path="/does/not/exist", loaded_workbook_cache_size=1)

    backend._openpyxl_workbook_pair(first)
    backend._openpyxl_workbook_pair(second)

    assert closed == ["first.xlsx:False", "first.xlsx:True"]


def test_libreoffice_backend_concurrent_prepared_cache_uses_one_conversion(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    cache_dir = tmp_path / "lo-cache"
    barrier = threading.Barrier(2)
    calls = 0
    errors: list[BaseException] = []
    results: list[dict[str, Any]] = []

    def fake_convert(
        _backend: LibreOfficeBackend,
        workbook_path: Path,
        temp_dir: Path,
        *,
        force_recalculate: bool,
    ) -> Path:
        nonlocal calls
        calls += 1
        time.sleep(0.05)
        converted = temp_dir / workbook_path.name
        shutil.copy2(workbook_path, converted)
        return converted

    def worker() -> None:
        try:
            barrier.wait(timeout=5)
            result = LibreOfficeBackend(
                soffice_path="/does/not/exist",
                cache_dir=cache_dir,
                timeout_seconds=5,
            ).recalculate(workbook)
            results.append(result)
        except BaseException as e:  # noqa: BLE001
            errors.append(e)

    monkeypatch.setattr(LibreOfficeBackend, "_convert_with_libreoffice", fake_convert)

    threads = [threading.Thread(target=worker) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert errors == []
    assert calls == 1
    assert len(results) == 2
    assert sorted(result["metrics"]["libreoffice_prepared_cache_hit"] for result in results) == [False, True]


def test_libreoffice_backend_forces_recalc_in_temp_profile(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)
    profile_xml: list[str] = []

    def fake_run(
        command: list[str],
        *,
        check: bool,
        capture_output: bool,
        text: bool,
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        assert check is False
        assert capture_output is True
        assert text is True
        assert timeout == 19
        profile_arg = next(arg for arg in command if arg.startswith("-env:UserInstallation="))
        profile_dir = Path(profile_arg.split("file://", 1)[1])
        profile_xml.append((profile_dir / "user" / "registrymodifications.xcu").read_text())
        output_dir = Path(command[command.index("--outdir") + 1])
        source_path = Path(command[-1])
        shutil.copy2(source_path, output_dir / f"{source_path.stem}.xlsx")
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr("gandalf.excel_mcp.subprocess.run", fake_run)

    result = LibreOfficeBackend(soffice_path="/usr/bin/soffice", timeout_seconds=19).recalculate(workbook)

    assert result["recalculated"] is True
    assert profile_xml
    assert '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>' in profile_xml[0]
    assert '<prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop>' in profile_xml[0]


def test_libreoffice_backend_repair_check_is_skipped(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    make_openpyxl_workbook(workbook)

    result = LibreOfficeBackend().workbook_repair_check(workbook)

    assert result["backend"] == "libreoffice"
    assert result["checked"] is False
    assert result["repair_dialog_detected"] is None
    assert result["skipped_reason"] == "repair dialogue detection requires Microsoft Excel"


def test_create_service_supports_libreoffice_backend(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class FakeLibreOfficeBackend(FakeBackend):
        def __init__(self, *, timeout_seconds: int, cache_dir: Path | None) -> None:
            self.timeout_seconds = timeout_seconds
            self.cache_dir = cache_dir

    monkeypatch.setattr(excel_mcp, "LibreOfficeBackend", FakeLibreOfficeBackend)
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"placeholder")
    cache_dir = tmp_path / "lo-cache"

    service = excel_mcp.create_service(
        ExcelServerConfig(
            workdir=tmp_path,
            backend="libreoffice",
            timeout_seconds=17,
            cache_dir=cache_dir,
        )
    )
    result = service.workbook_summary("model.xlsx", recalculate=False)

    assert result["backend"] == "fake"
    assert isinstance(service.backend, FakeLibreOfficeBackend)
    assert service.backend.timeout_seconds == 17
    assert service.backend.cache_dir == cache_dir


def test_config_from_env_reads_libreoffice_cache_dir(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    cache_dir = tmp_path / "lo-cache"
    monkeypatch.setenv("GANDALF_EXCEL_WORKDIR", str(tmp_path))
    monkeypatch.setenv("GANDALF_EXCEL_BACKEND", "libreoffice")
    monkeypatch.setenv("GANDALF_EXCEL_CACHE_DIR", str(cache_dir))

    config = config_from_env()

    assert config.backend == "libreoffice"
    assert config.cache_dir == cache_dir


def test_range_helpers() -> None:
    assert range_cell_count("A1") == 1
    assert range_cell_count("A1:C3") == 9
    assert cell_address(12, 28) == "AB12"
    with pytest.raises(ValueError, match="Invalid"):
        range_cell_count("B2:A1")
    with pytest.raises(ValueError, match="maximum Excel column XFD"):
        range_cell_count("XFE1")


def test_range_number_formats_returns_per_cell_grid() -> None:
    class FakeCell:
        def __init__(self, row: int, column: int) -> None:
            self.number_format = f"fmt-{row}-{column}"

    class FakeWorksheet:
        def range(self, address: tuple[int, int]) -> FakeCell:
            row, column = address
            return FakeCell(row, column)

    class FakeRange:
        row = 3
        column = 2

    formats = range_number_formats(FakeWorksheet(), FakeRange(), rows=2, columns=2)

    assert formats == [["fmt-3-2", "fmt-3-3"], ["fmt-4-2", "fmt-4-3"]]


def test_config_from_env(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv("GANDALF_EXCEL_WORKDIR", str(tmp_path))
    monkeypatch.setenv("GANDALF_EXCEL_BACKEND", "windows_vm")
    monkeypatch.setenv("GANDALF_EXCEL_WINDOWS_URL", "http://excel-vm")
    monkeypatch.setenv("GANDALF_EXCEL_TIMEOUT_SECONDS", "45")
    monkeypatch.setenv("GANDALF_EXCEL_VISIBLE", "true")
    monkeypatch.setenv("GANDALF_EXCEL_ALLOW_MACROS", "yes")
    monkeypatch.setenv("GANDALF_EXCEL_MAX_CELLS_PER_CALL", "456")
    monkeypatch.setenv("GANDALF_EXCEL_MAX_FORMAT_CELLS_PER_CALL", "123")
    monkeypatch.setenv("GANDALF_EXCEL_AUTH_TOKEN", "secret")

    config = config_from_env()

    assert config.workdir == tmp_path
    assert config.backend == "windows_vm"
    assert config.windows_url == "http://excel-vm"
    assert config.timeout_seconds == 45
    assert config.visible is True
    assert config.allow_macros is True
    assert config.max_cells_per_call == 456
    assert config.max_format_cells_per_call == 123
    assert config.auth_token == "secret"


def test_windows_vm_proxy_posts_workbook_payload(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook bytes")
    captured: list[dict[str, Any]] = []

    class FakeResponse:
        def __init__(self, payload: dict[str, Any]) -> None:
            self.payload = payload

        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self) -> bytes:
            return json.dumps(self.payload).encode()

    def fake_urlopen(request: Any, *, timeout: int) -> FakeResponse:
        payload = json.loads(request.data.decode("utf-8"))
        captured.append(
            {
                "timeout": timeout,
                "url": request.full_url,
                "headers": {key.lower(): value for key, value in request.header_items()},
                "payload": payload,
            }
        )
        if request.full_url.endswith("/workbook/ensure") and "workbook_b64" not in payload:
            return FakeResponse({"ok": False, "error": "workbook_not_cached"})
        if request.full_url.endswith("/workbook/ensure"):
            return FakeResponse(
                {
                    "ok": True,
                    "result": {
                        "workbook_id": payload["workbook_sha256"],
                        "workbook_name": payload["workbook_name"],
                        "cache_hit": False,
                    },
                }
            )
        return FakeResponse({"ok": True, "result": {"path": "model.xlsx", "backend": "windows_vm"}})

    backend = WindowsVMExcelBackend(url="http://excel-vm.local", timeout_seconds=12, auth_token="secret")
    with patch("gandalf.excel_mcp.urllib.request.urlopen", side_effect=fake_urlopen):
        result = backend.inspect_range(
            workbook,
            sheet="Sheet1",
            range_address="A1",
            include_values=True,
            include_formulas=True,
            include_formats=False,
            recalculate=True,
        )
        backend.find_cells(
            workbook,
            query="Revenue",
            search_values=True,
            search_formulas=True,
            max_results=1,
            max_cells_per_call=10,
        )

    assert result["path"] == "model.xlsx"
    assert result["backend"] == "windows_vm"
    assert result["metrics"]["windows_vm_client_workbook_bytes_sent"] == len(b"workbook bytes")
    assert [call["url"] for call in captured] == [
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/inspect",
        "http://excel-vm.local/inspect",
    ]
    assert all(call["timeout"] == 12 for call in captured)
    assert all(call["headers"]["x-gandalf-excel-token"] == "secret" for call in captured)
    assert "workbook_b64" not in captured[0]["payload"]
    assert captured[1]["payload"]["workbook_b64"]
    assert captured[2]["payload"]["operation"] == "inspect_range"
    assert captured[2]["payload"]["workbook_name"] == "model.xlsx"
    assert captured[2]["payload"]["args"]["range_address"] == "A1"
    assert captured[2]["payload"]["workbook_id"] == captured[1]["payload"]["workbook_sha256"]
    assert "workbook_b64" not in captured[2]["payload"]
    assert captured[3]["payload"]["operation"] == "find_cells"
    assert "workbook_b64" not in captured[3]["payload"]


def test_windows_vm_proxy_falls_back_for_legacy_worker_without_ensure_endpoint(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook bytes")
    captured: list[dict[str, Any]] = []

    class FakeResponse:
        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self) -> bytes:
            return json.dumps({"ok": True, "result": {"path": "model.xlsx", "backend": "windows_vm"}}).encode()

    def fake_urlopen(request: Any, *, timeout: int) -> FakeResponse:
        payload = json.loads(request.data.decode("utf-8"))
        captured.append({"timeout": timeout, "url": request.full_url, "payload": payload})
        if request.full_url.endswith("/workbook/ensure"):
            raise urllib.error.HTTPError(
                request.full_url,
                404,
                "Not Found",
                hdrs=Message(),
                fp=None,
            )
        return FakeResponse()

    backend = WindowsVMExcelBackend(url="http://excel-vm.local", timeout_seconds=12)
    with patch("gandalf.excel_mcp.urllib.request.urlopen", side_effect=fake_urlopen):
        result = backend.workbook_summary(workbook, recalculate=False)
        backend.find_cells(
            workbook,
            query="Revenue",
            search_values=True,
            search_formulas=True,
            max_results=1,
            max_cells_per_call=10,
        )

    assert result["backend"] == "windows_vm"
    assert result["metrics"]["windows_vm_client_legacy_full_upload"] is True
    assert result["metrics"]["windows_vm_client_workbook_bytes_sent"] == len(b"workbook bytes")
    assert [call["url"] for call in captured] == [
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/inspect",
        "http://excel-vm.local/inspect",
    ]
    assert captured[1]["payload"]["operation"] == "workbook_summary"
    assert captured[1]["payload"]["workbook_b64"]
    assert captured[2]["payload"]["operation"] == "find_cells"
    assert captured[2]["payload"]["workbook_b64"]


def test_windows_vm_proxy_recovers_when_worker_cache_is_evicted_before_inspect(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook bytes")
    captured: list[dict[str, Any]] = []
    find_inspect_attempts = 0

    class FakeResponse:
        def __init__(self, payload: dict[str, Any]) -> None:
            self.payload = payload

        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self) -> bytes:
            return json.dumps(self.payload).encode()

    def fake_urlopen(request: Any, *, timeout: int) -> FakeResponse:
        nonlocal find_inspect_attempts
        payload = json.loads(request.data.decode("utf-8"))
        captured.append({"timeout": timeout, "url": request.full_url, "payload": payload})
        if request.full_url.endswith("/workbook/ensure") and "workbook_b64" not in payload:
            return FakeResponse({"ok": False, "error": "workbook_not_cached"})
        if request.full_url.endswith("/workbook/ensure"):
            return FakeResponse(
                {
                    "ok": True,
                    "result": {
                        "workbook_id": payload["workbook_sha256"],
                        "workbook_name": payload["workbook_name"],
                        "cache_hit": False,
                    },
                }
            )
        if payload["operation"] == "find_cells" and find_inspect_attempts == 0:
            find_inspect_attempts += 1
            return FakeResponse({"ok": False, "error": "workbook_not_cached"})
        return FakeResponse({"ok": True, "result": {"path": "model.xlsx", "backend": "windows_vm"}})

    backend = WindowsVMExcelBackend(url="http://excel-vm.local", timeout_seconds=12)
    with patch("gandalf.excel_mcp.urllib.request.urlopen", side_effect=fake_urlopen):
        backend.workbook_summary(workbook, recalculate=False)
        result = backend.find_cells(
            workbook,
            query="Revenue",
            search_values=True,
            search_formulas=True,
            max_results=1,
            max_cells_per_call=10,
        )

    assert result["backend"] == "windows_vm"
    assert result["metrics"]["windows_vm_client_retried_after_cache_miss"] is True
    assert result["metrics"]["windows_vm_client_workbook_bytes_sent"] == len(b"workbook bytes")
    assert [call["url"] for call in captured] == [
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/inspect",
        "http://excel-vm.local/inspect",
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/workbook/ensure",
        "http://excel-vm.local/inspect",
    ]
    assert captured[3]["payload"]["operation"] == "find_cells"
    assert "workbook_b64" not in captured[3]["payload"]
    assert captured[5]["payload"]["workbook_b64"]
    assert captured[6]["payload"]["operation"] == "find_cells"


def test_windows_vm_proxy_posts_repair_check(tmp_path: Path) -> None:
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook bytes")
    captured: list[dict[str, Any]] = []

    class FakeResponse:
        def __init__(self, payload: dict[str, Any]) -> None:
            self.payload = payload

        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self) -> bytes:
            return json.dumps(self.payload).encode()

    def fake_urlopen(request: Any, *, timeout: int) -> FakeResponse:
        payload = json.loads(request.data.decode("utf-8"))
        captured.append({"timeout": timeout, "url": request.full_url, "payload": payload})
        if request.full_url.endswith("/workbook/ensure") and "workbook_b64" not in payload:
            return FakeResponse({"ok": False, "error": "workbook_not_cached"})
        if request.full_url.endswith("/workbook/ensure"):
            return FakeResponse(
                {
                    "ok": True,
                    "result": {
                        "workbook_id": payload["workbook_sha256"],
                        "workbook_name": payload["workbook_name"],
                        "cache_hit": False,
                    },
                }
            )
        return FakeResponse(
            {
                "ok": True,
                "result": {
                    "path": "model.xlsx",
                    "backend": "windows_vm",
                    "checked": True,
                    "opened": True,
                    "repair_dialog_detected": True,
                },
            }
        )

    backend = WindowsVMExcelBackend(url="http://excel-vm.local", timeout_seconds=12)
    with patch("gandalf.excel_mcp.urllib.request.urlopen", side_effect=fake_urlopen):
        result = backend.workbook_repair_check(workbook)

    assert result["repair_dialog_detected"] is True
    assert result["backend"] == "windows_vm"
    assert all(call["timeout"] == 12 for call in captured)
    assert captured[-1]["url"] == "http://excel-vm.local/inspect"
    assert captured[-1]["payload"]["operation"] == "workbook_repair_check"
    assert "workbook_b64" not in captured[-1]["payload"]


def test_local_excel_repair_check_detects_repair_probe_success(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook bytes")
    calls: list[int] = []

    def fake_probe_open(_backend: excel_mcp.LocalExcelBackend, _path: Path, *, corrupt_load: int) -> dict[str, Any]:
        calls.append(corrupt_load)
        if corrupt_load == excel_mcp._XL_NORMAL_LOAD:
            return {
                "opened": False,
                "corrupt_load": corrupt_load,
                "error": "Excel found unreadable content",
            }
        return {
            "opened": True,
            "corrupt_load": corrupt_load,
            "sheet_count": 1,
        }

    monkeypatch.setattr(excel_mcp.LocalExcelBackend, "_probe_open", fake_probe_open)

    result = excel_mcp.LocalExcelBackend().workbook_repair_check(workbook)

    assert result["repair_dialog_detected"] is True
    assert result["opened"] is True
    assert calls == [excel_mcp._XL_NORMAL_LOAD, excel_mcp._XL_REPAIR_FILE]


def test_corrupt_load_unsupported_is_not_repair_evidence() -> None:
    error = "corrupt_load is not supported on macOS"

    assert excel_mcp.corrupt_load_is_unsupported(error) is True
    assert excel_mcp.error_mentions_repair(error) is False


def test_worker_request_auth_helper() -> None:
    assert worker_request_is_authorised(None, None) is True
    assert worker_request_is_authorised("secret", "secret") is True
    assert worker_request_is_authorised("secret", None) is False
    assert worker_request_is_authorised("secret", "wrong") is False
    assert worker_request_is_authorised("secrét", "secrét") is True
    assert worker_request_is_authorised("secrét", "secret") is False


def test_prune_worker_cache_removes_only_stale_session_dirs(tmp_path: Path) -> None:
    old_dir = tmp_path / "old"
    new_dir = tmp_path / "new"
    old_dir.mkdir()
    new_dir.mkdir()
    marker = tmp_path / "not-a-session"
    marker.write_text("keep")
    now = time.time()
    os.utime(old_dir, (now - 1_000, now - 1_000))
    os.utime(new_dir, (now, now))

    prune_worker_cache(tmp_path, ttl_seconds=100, now=now)

    assert not old_dir.exists()
    assert new_dir.exists()
    assert marker.exists()


def test_prune_worker_cache_skips_locked_session_dirs(tmp_path: Path) -> None:
    old_dir = tmp_path / "old"
    old_dir.mkdir()
    now = time.time()
    os.utime(old_dir, (now - 1_000, now - 1_000))
    with worker_cache_lock(old_dir.name):
        prune_worker_cache(tmp_path, ttl_seconds=100, now=now)

    assert old_dir.exists()


def test_worker_cache_lock_registry_releases_unused_entries() -> None:
    cache_key = "temporary-cache-key"

    with worker_cache_lock(cache_key) as acquired:
        assert acquired is True
        assert cache_key in excel_mcp._WORKER_CACHE_LOCKS

    assert cache_key not in excel_mcp._WORKER_CACHE_LOCKS


def worker_payload(
    operation: str,
    args: dict[str, Any] | None = None,
    *,
    workbook_bytes: bytes = b"workbook",
) -> dict[str, Any]:
    return {
        "operation": operation,
        "args": args or {},
        "workbook_name": "model.xlsx",
        "workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "workbook_b64": base64.b64encode(workbook_bytes).decode("ascii"),
    }


def test_handle_worker_ensure_uploads_once_then_accepts_cached_reference(tmp_path: Path) -> None:
    payload = worker_payload("workbook_summary")
    ensure_payload = {
        "workbook_name": payload["workbook_name"],
        "workbook_sha256": payload["workbook_sha256"],
        "workbook_b64": payload["workbook_b64"],
    }

    first = handle_worker_ensure(ensure_payload, temp_dir=tmp_path)
    second = handle_worker_ensure(
        {
            "workbook_name": payload["workbook_name"],
            "workbook_sha256": payload["workbook_sha256"],
            "workbook_id": payload["workbook_sha256"],
        },
        temp_dir=tmp_path,
    )

    assert first["workbook_id"] == payload["workbook_sha256"]
    assert first["cache_hit"] is False
    assert first["uploaded_bytes"] == len(b"workbook")
    assert second["cache_hit"] is True
    assert second["uploaded_bytes"] == 0


def test_handle_worker_ensure_reports_cache_miss_without_upload(tmp_path: Path) -> None:
    payload = worker_payload("workbook_summary")

    with pytest.raises(FileNotFoundError, match="workbook_not_cached"):
        handle_worker_ensure(
            {
                "workbook_name": payload["workbook_name"],
                "workbook_sha256": payload["workbook_sha256"],
                "workbook_id": payload["workbook_sha256"],
            },
            temp_dir=tmp_path,
        )


def test_handle_worker_request_rejects_sha_mismatch(tmp_path: Path) -> None:
    payload = worker_payload("workbook_summary")
    payload["workbook_sha256"] = "not-the-real-sha"

    with pytest.raises(ValueError, match="sha256"):
        handle_worker_request(
            payload,
            temp_dir=tmp_path,
            visible=False,
            allow_macros=False,
            max_cells_per_call=10,
        )


def test_handle_worker_request_rejects_unknown_operation(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unsupported operation"):
        handle_worker_request(
            worker_payload("unknown_operation"),
            temp_dir=tmp_path,
            visible=False,
            allow_macros=False,
            max_cells_per_call=10,
        )


def test_handle_worker_request_rejects_zero_cache_ttl(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="ttl_seconds must be positive"):
        handle_worker_request(
            worker_payload("workbook_summary"),
            temp_dir=tmp_path,
            visible=False,
            allow_macros=False,
            max_cells_per_call=10,
            cache_ttl_seconds=0,
        )


def test_handle_worker_request_dispatches_to_backend(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class WorkerFakeBackend(FakeBackend):
        def __init__(self, *, visible: bool) -> None:
            self.visible = visible

    monkeypatch.setattr("gandalf.excel_mcp.LocalExcelBackend", WorkerFakeBackend)

    result = handle_worker_request(
        worker_payload("workbook_summary", {"recalculate": False}),
        temp_dir=tmp_path,
        visible=True,
        allow_macros=False,
        max_cells_per_call=10,
    )

    assert result["backend"] == "fake"
    assert result["recalculated"] is False
    assert result["metrics"]["windows_vm_worker_workbook_cache_hit"] is False


def test_handle_worker_request_accepts_cached_workbook_reference(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class WorkerFakeBackend(FakeBackend):
        def __init__(self, *, visible: bool) -> None:
            self.visible = visible

    monkeypatch.setattr("gandalf.excel_mcp.LocalExcelBackend", WorkerFakeBackend)
    payload = worker_payload("workbook_summary", {"recalculate": False})
    handle_worker_ensure(
        {
            "workbook_name": payload["workbook_name"],
            "workbook_sha256": payload["workbook_sha256"],
            "workbook_b64": payload["workbook_b64"],
        },
        temp_dir=tmp_path,
    )

    result = handle_worker_request(
        {
            "operation": "workbook_summary",
            "args": {"recalculate": False},
            "workbook_name": payload["workbook_name"],
            "workbook_sha256": payload["workbook_sha256"],
            "workbook_id": payload["workbook_sha256"],
        },
        temp_dir=tmp_path,
        visible=True,
        allow_macros=False,
        max_cells_per_call=10,
    )

    assert result["backend"] == "fake"
    assert result["recalculated"] is False
    assert result["metrics"]["windows_vm_worker_workbook_cache_hit"] is True


def test_handle_worker_request_dispatches_repair_check(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class WorkerFakeBackend(FakeBackend):
        def __init__(self, *, visible: bool) -> None:
            self.visible = visible

    monkeypatch.setattr("gandalf.excel_mcp.LocalExcelBackend", WorkerFakeBackend)

    result = handle_worker_request(
        worker_payload("workbook_repair_check"),
        temp_dir=tmp_path,
        visible=True,
        allow_macros=False,
        max_cells_per_call=10,
    )

    assert result["backend"] == "fake"
    assert result["repair_dialog_detected"] is False


def test_handle_worker_request_refreshes_cache_hit_mtime(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class WorkerFakeBackend(FakeBackend):
        def __init__(self, *, visible: bool) -> None:
            self.visible = visible

    monkeypatch.setattr("gandalf.excel_mcp.LocalExcelBackend", WorkerFakeBackend)
    payload = worker_payload("workbook_summary", {"recalculate": False})
    actual_sha = payload["workbook_sha256"]
    handle_worker_request(
        payload,
        temp_dir=tmp_path,
        visible=False,
        allow_macros=False,
        max_cells_per_call=10,
    )
    session_dir = tmp_path / actual_sha
    workbook_path = session_dir / "model.xlsx"
    old_time = time.time() - 1_000
    os.utime(session_dir, (old_time, old_time))
    os.utime(workbook_path, (old_time, old_time))

    handle_worker_request(
        payload,
        temp_dir=tmp_path,
        visible=False,
        allow_macros=False,
        max_cells_per_call=10,
    )

    assert session_dir.stat().st_mtime > old_time
    assert workbook_path.stat().st_mtime > old_time


def test_handle_worker_request_applies_format_cell_limit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class WorkerFakeBackend(FakeBackend):
        def __init__(self, *, visible: bool) -> None:
            self.visible = visible

    monkeypatch.setattr("gandalf.excel_mcp.LocalExcelBackend", WorkerFakeBackend)

    with pytest.raises(ValueError, match="max_format_cells_per_call"):
        handle_worker_request(
            worker_payload(
                "inspect_range",
                {
                    "sheet": "Sheet1",
                    "range_address": "A1:B1",
                    "include_formats": True,
                },
            ),
            temp_dir=tmp_path,
            visible=False,
            allow_macros=False,
            max_cells_per_call=10,
            max_format_cells_per_call=1,
        )


def test_worker_launch_validates_cache_ttl(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="cache_ttl_seconds must be positive"):
        run_worker(
            host="127.0.0.1",
            port=8765,
            temp_dir=tmp_path,
            auth_token=None,
            visible=False,
            allow_macros=False,
            max_cells_per_call=10,
            cache_ttl_seconds=0,
        )

    with pytest.raises(SystemExit):
        worker_main(["--cache-ttl-seconds", "0"])


@pytest.mark.excel
def test_mac_excel_backend_integration_is_opt_in(tmp_path: Path) -> None:
    if os.environ.get("GANDALF_RUN_EXCEL_INTEGRATION") != "1":
        pytest.skip("Set GANDALF_RUN_EXCEL_INTEGRATION=1 to run Microsoft Excel integration tests")
    excel_app = Path("/Applications/Microsoft Excel.app")
    if not excel_app.exists():
        pytest.skip("Microsoft Excel.app is not installed")
    pytest.importorskip("xlwings")

    # The real integration workbook is intentionally left to opt-in runs
    # because macOS prompts for Excel automation permissions on first use.
    assert tmp_path.exists()


class _StubLocalBackend:
    """Stub for the worker's local Excel backend (no real Excel needed)."""

    def __init__(self, *, visible: bool = False) -> None:
        self.visible = visible

    def workbook_summary(self, workbook_path: Path, *, recalculate: bool) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "mac_excel",
            "recalculated": recalculate,
            "sheets": [{"name": "Sheet1"}],
        }

    def inspect_range(self, workbook_path: Path, *, sheet: str, range_address: str, **_kwargs: Any) -> dict[str, Any]:
        return {"path": workbook_path.name, "backend": "mac_excel", "sheet": sheet, "range": range_address}

    def find_cells(self, workbook_path: Path, *, query: str, **_kwargs: Any) -> dict[str, Any]:
        return {"path": workbook_path.name, "backend": "mac_excel", "query": query, "matches": []}

    def recalculate(self, workbook_path: Path) -> dict[str, Any]:
        return {"path": workbook_path.name, "backend": "mac_excel", "recalculated": True}

    def workbook_repair_check(self, workbook_path: Path) -> dict[str, Any]:
        return {
            "path": workbook_path.name,
            "backend": "mac_excel",
            "checked": True,
            "opened": True,
            "repair_dialog_detected": False,
        }


@contextmanager
def _running_worker(tmp_path: Path, *, auth_token: str | None = None) -> Iterator[ThreadingHTTPServer]:
    """Run the worker HTTP server on an ephemeral port for the duration of the block."""
    server = build_worker_server(
        host="127.0.0.1",
        port=0,
        temp_dir=tmp_path / "worker",
        auth_token=auth_token,
        visible=False,
        allow_macros=False,
        max_cells_per_call=10_000,
    )
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield server
    finally:
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()


def _worker_url(server: ThreadingHTTPServer) -> str:
    return f"http://127.0.0.1:{server.server_address[1]}"


def test_windows_vm_round_trip_summary(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel_mcp, "LocalExcelBackend", _StubLocalBackend)
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook-bytes")

    with _running_worker(tmp_path) as server:
        backend = WindowsVMExcelBackend(url=_worker_url(server), timeout_seconds=5)
        result = backend.workbook_summary(workbook, recalculate=False)

    # Real HTTP round trip: proxy -> worker -> stub backend -> proxy.
    assert result["backend"] == "windows_vm"  # proxy stamps the backend
    assert result["sheets"] == [{"name": "Sheet1"}]
    assert result["recalculated"] is False


def test_windows_vm_round_trip_repair_check(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel_mcp, "LocalExcelBackend", _StubLocalBackend)
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook-bytes")

    with _running_worker(tmp_path) as server:
        backend = WindowsVMExcelBackend(url=_worker_url(server), timeout_seconds=5)
        result = backend.workbook_repair_check(workbook)

    assert result["backend"] == "windows_vm"
    assert result["repair_dialog_detected"] is False
    assert result["opened"] is True


def test_windows_vm_round_trip_enforces_auth_token(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel_mcp, "LocalExcelBackend", _StubLocalBackend)
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook-bytes")

    with _running_worker(tmp_path, auth_token="s3cret") as server:
        url = _worker_url(server)
        missing_token = WindowsVMExcelBackend(url=url, timeout_seconds=5)
        # A 401 surfaces through urllib as an HTTPError (request-failed path).
        with pytest.raises(RuntimeError, match="401"):
            missing_token.workbook_summary(workbook, recalculate=False)

        with_token = WindowsVMExcelBackend(url=url, timeout_seconds=5, auth_token="s3cret")
        result = with_token.workbook_summary(workbook, recalculate=False)

    assert result["backend"] == "windows_vm"


def test_windows_vm_round_trip_surfaces_worker_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    class FailingBackend(_StubLocalBackend):
        def workbook_summary(self, workbook_path: Path, *, recalculate: bool) -> dict[str, Any]:
            msg = "stub failure"
            raise ValueError(msg)

    monkeypatch.setattr(excel_mcp, "LocalExcelBackend", FailingBackend)
    workbook = tmp_path / "model.xlsx"
    workbook.write_bytes(b"workbook-bytes")

    with _running_worker(tmp_path) as server:
        backend = WindowsVMExcelBackend(url=_worker_url(server), timeout_seconds=5)
        with pytest.raises(RuntimeError, match="stub failure"):
            backend.workbook_summary(workbook, recalculate=False)
