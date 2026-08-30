"""Tests for orchestrator-level functions (resolve_judge_guidance, evaluate_all_criteria)."""

import json
import os
import pathlib
import shutil
import subprocess
import tempfile
import threading
import time
from collections.abc import Callable
from typing import Any
from unittest.mock import patch

import pytest

from gandalf.models import (
    BatchJudgeInput,
    CriterionResult,
    ExcelBackendConfig,
    GraderConfig,
    JudgeInput,
    LLMUsage,
    RubricItem,
    SectionGateResult,
    Verdict,
    WorkbookRepairCheckResult,
)
from gandalf.orchestrator import (
    JUDGE_ENV_ALLOWLIST,
    apply_golden_check_defaults,
    apply_excel_backend_cli_override,
    clone_workspace,
    collect_excel_metrics,
    collect_section_gates,
    criterion_for_judge,
    effective_judge_guidance,
    evaluate_section_gates,
    filter_rubric_sections,
    default_repair_check_workbook,
    judge_clone_parent_dir,
    judge_env_vars,
    libreoffice_cache_environment,
    main,
    resolve_instructions,
    resolve_judge_guidance,
    resolve_judge_prompt,
    run_excel_preflight,
    run_workbook_repair_check,
    run_batch,
    run_batch_concurrent,
    warn_if_excel_workspace_protected,
    workbook_digest_text,
    run_individual,
    run_judge,
    write_info,
)
from openpyxl import Workbook  # type: ignore[import-untyped]

from tests.conftest import cr, make_batch_input, make_config


class TestResolveInstructions:
    def test_no_config_or_env_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_INSTRUCTIONS_PATH", raising=False)
        config = make_config(instructions=None)
        with pytest.raises(SystemExit):
            resolve_instructions(config)

    def test_inline_returns_content(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_INSTRUCTIONS_PATH", raising=False)
        config = make_config(instructions="inline instructions")
        assert resolve_instructions(config) == "inline instructions"

    def test_reads_file_from_toml_path(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_INSTRUCTIONS_PATH", raising=False)
        instructions_file = tmp_path / "instructions.md"
        instructions_file.write_text("Instructions from file.")
        config = make_config(instructions=None, instructions_path=str(instructions_file))
        assert resolve_instructions(config) == "Instructions from file."

    def test_reads_file_from_env_var(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        instructions_file = tmp_path / "instructions.md"
        instructions_file.write_text("From env var.")
        monkeypatch.setenv("GRADER_INSTRUCTIONS_PATH", str(instructions_file))
        config = make_config(instructions=None)
        assert resolve_instructions(config) == "From env var."

    def test_toml_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        toml_file = tmp_path / "toml_instructions.md"
        toml_file.write_text("From TOML.")
        env_file = tmp_path / "env_instructions.md"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_INSTRUCTIONS_PATH", str(env_file))
        config = make_config(instructions=None, instructions_path=str(toml_file))
        assert resolve_instructions(config) == "From TOML."

    def test_inline_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        env_file = tmp_path / "env_instructions.md"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_INSTRUCTIONS_PATH", str(env_file))
        config = make_config(instructions="inline wins")
        assert resolve_instructions(config) == "inline wins"

    def test_missing_configured_toml_path_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_INSTRUCTIONS_PATH", raising=False)
        config = make_config(instructions=None, instructions_path="/nonexistent/instructions.md")
        with pytest.raises(SystemExit):
            resolve_instructions(config)

    def test_missing_configured_env_path_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("GRADER_INSTRUCTIONS_PATH", "/nonexistent/instructions.md")
        config = make_config(instructions=None)
        with pytest.raises(SystemExit):
            resolve_instructions(config)

    def test_error_message_mentions_file_path(
        self, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.delenv("GRADER_INSTRUCTIONS_PATH", raising=False)
        config = make_config(instructions=None, instructions_path="/missing/instructions.md")
        with pytest.raises(SystemExit):
            resolve_instructions(config)
        stderr = capsys.readouterr().err
        assert "/missing/instructions.md" in stderr
        assert "instructions_path" in stderr

    def test_error_message_mentions_env_var_source(
        self, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setenv("GRADER_INSTRUCTIONS_PATH", "/missing/env_instructions.md")
        config = make_config(instructions=None)
        with pytest.raises(SystemExit):
            resolve_instructions(config)
        stderr = capsys.readouterr().err
        assert "/missing/env_instructions.md" in stderr
        assert "GRADER_INSTRUCTIONS_PATH" in stderr


class TestResolveJudgeGuidance:
    def test_no_path_returns_empty(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_GUIDANCE_PATH", raising=False)
        config = make_config()
        assert resolve_judge_guidance(config) == ""

    def test_reads_file_from_toml_path(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_GUIDANCE_PATH", raising=False)
        guidance_file = tmp_path / "guidance.md"
        guidance_file.write_text("Use openpyxl for .xlsx files.")
        config = make_config(judge_guidance_path=str(guidance_file))
        assert resolve_judge_guidance(config) == "Use openpyxl for .xlsx files."

    def test_reads_file_from_env_var(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        guidance_file = tmp_path / "guidance.md"
        guidance_file.write_text("From env var.")
        monkeypatch.setenv("GRADER_JUDGE_GUIDANCE_PATH", str(guidance_file))
        config = make_config()  # no judge_guidance_path in TOML
        assert resolve_judge_guidance(config) == "From env var."

    def test_toml_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        toml_file = tmp_path / "toml_guidance.md"
        toml_file.write_text("From TOML.")
        env_file = tmp_path / "env_guidance.md"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_JUDGE_GUIDANCE_PATH", str(env_file))
        config = make_config(judge_guidance_path=str(toml_file))
        assert resolve_judge_guidance(config) == "From TOML."

    def test_missing_configured_toml_path_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_GUIDANCE_PATH", raising=False)
        config = make_config(judge_guidance_path="/nonexistent/guidance.md")
        with pytest.raises(SystemExit):
            resolve_judge_guidance(config)

    def test_missing_configured_env_path_exits(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:  # noqa: ARG002
        monkeypatch.setenv("GRADER_JUDGE_GUIDANCE_PATH", "/nonexistent/guidance.md")
        config = make_config()
        with pytest.raises(SystemExit):
            resolve_judge_guidance(config)

    def test_error_message_mentions_file_path(
        self, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.delenv("GRADER_JUDGE_GUIDANCE_PATH", raising=False)
        config = make_config(judge_guidance_path="/missing/guidance.md")
        with pytest.raises(SystemExit):
            resolve_judge_guidance(config)
        stderr = capsys.readouterr().err
        assert "/missing/guidance.md" in stderr
        assert "judge_guidance_path" in stderr

    def test_error_message_mentions_env_var_source(
        self, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setenv("GRADER_JUDGE_GUIDANCE_PATH", "/missing/env_guidance.md")
        config = make_config()
        with pytest.raises(SystemExit):
            resolve_judge_guidance(config)
        stderr = capsys.readouterr().err
        assert "/missing/env_guidance.md" in stderr
        assert "GRADER_JUDGE_GUIDANCE_PATH" in stderr


class TestResolveJudgeGuidanceInline:
    """Tests for inline judge_guidance (no file path)."""

    def test_inline_returns_content(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_GUIDANCE_PATH", raising=False)
        config = make_config(judge_guidance="inline guidance text")
        assert resolve_judge_guidance(config) == "inline guidance text"

    def test_inline_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        env_file = tmp_path / "env_guidance.md"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_JUDGE_GUIDANCE_PATH", str(env_file))
        config = make_config(judge_guidance="inline wins")
        assert resolve_judge_guidance(config) == "inline wins"


class TestResolveJudgePrompt:
    def test_no_config_returns_none(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_PROMPT_PATH", raising=False)
        config = make_config()
        assert resolve_judge_prompt(config) is None

    def test_inline_returns_content(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_PROMPT_PATH", raising=False)
        config = make_config(judge_prompt="Hello {{ instructions }}")
        assert resolve_judge_prompt(config) == "Hello {{ instructions }}"

    def test_reads_file_from_path(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_PROMPT_PATH", raising=False)
        template_file = tmp_path / "prompt.j2"
        template_file.write_text("Custom {{ criterion }}")
        config = make_config(judge_prompt_path=str(template_file))
        assert resolve_judge_prompt(config) == "Custom {{ criterion }}"

    def test_reads_file_from_env_var(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        template_file = tmp_path / "prompt.j2"
        template_file.write_text("From env var.")
        monkeypatch.setenv("GRADER_JUDGE_PROMPT_PATH", str(template_file))
        config = make_config()
        assert resolve_judge_prompt(config) == "From env var."

    def test_path_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        toml_file = tmp_path / "toml_prompt.j2"
        toml_file.write_text("From TOML.")
        env_file = tmp_path / "env_prompt.j2"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_JUDGE_PROMPT_PATH", str(env_file))
        config = make_config(judge_prompt_path=str(toml_file))
        assert resolve_judge_prompt(config) == "From TOML."

    def test_inline_takes_precedence_over_env(self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
        env_file = tmp_path / "env_prompt.j2"
        env_file.write_text("From env.")
        monkeypatch.setenv("GRADER_JUDGE_PROMPT_PATH", str(env_file))
        config = make_config(judge_prompt="inline wins")
        assert resolve_judge_prompt(config) == "inline wins"

    def test_missing_path_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv("GRADER_JUDGE_PROMPT_PATH", raising=False)
        config = make_config(judge_prompt_path="/nonexistent/prompt.j2")
        with pytest.raises(SystemExit):
            resolve_judge_prompt(config)

    def test_missing_env_path_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("GRADER_JUDGE_PROMPT_PATH", "/nonexistent/prompt.j2")
        config = make_config()
        with pytest.raises(SystemExit):
            resolve_judge_prompt(config)

    def test_error_message_mentions_path(
        self,
        capsys: pytest.CaptureFixture[str],
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.delenv("GRADER_JUDGE_PROMPT_PATH", raising=False)
        config = make_config(judge_prompt_path="/missing/prompt.j2")
        with pytest.raises(SystemExit):
            resolve_judge_prompt(config)
        stderr = capsys.readouterr().err
        assert "/missing/prompt.j2" in stderr
        assert "judge_prompt_path" in stderr

    def test_error_message_mentions_env_var(
        self,
        capsys: pytest.CaptureFixture[str],
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.setenv("GRADER_JUDGE_PROMPT_PATH", "/missing/env_prompt.j2")
        config = make_config()
        with pytest.raises(SystemExit):
            resolve_judge_prompt(config)
        stderr = capsys.readouterr().err
        assert "/missing/env_prompt.j2" in stderr
        assert "GRADER_JUDGE_PROMPT_PATH" in stderr


class TestJudgeEnvVars:
    """Tests for the env-var allowlist forwarded to the judge subprocess."""

    def test_only_allowlisted_vars_are_forwarded(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("LLM_API_KEY", "sk-test-123")
        monkeypatch.setenv("PATH", "/usr/bin")
        monkeypatch.setenv("SECRET_TOKEN", "should-not-leak")
        monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "should-not-leak")
        result = judge_env_vars()
        keys = {item.split("=", 1)[0] for item in result}
        assert "LLM_API_KEY" in keys
        assert "PATH" in keys
        assert "SECRET_TOKEN" not in keys
        assert "AWS_SECRET_ACCESS_KEY" not in keys

    def test_empty_values_are_skipped(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setenv("LLM_API_KEY", "")
        monkeypatch.setenv("LLM_BASE_URL", "https://api.example.com")
        result = judge_env_vars()
        keys = {item.split("=", 1)[0] for item in result}
        assert "LLM_API_KEY" not in keys
        assert "LLM_BASE_URL" in keys

    def test_missing_vars_are_silently_skipped(self, monkeypatch: pytest.MonkeyPatch) -> None:
        for key in JUDGE_ENV_ALLOWLIST:
            monkeypatch.delenv(key, raising=False)
        assert judge_env_vars() == []

    def test_all_allowlisted_vars_forwarded_when_present(self, monkeypatch: pytest.MonkeyPatch) -> None:
        for key in JUDGE_ENV_ALLOWLIST:
            monkeypatch.setenv(key, f"val-{key}")
        result = judge_env_vars()
        keys = {item.split("=", 1)[0] for item in result}
        assert keys == JUDGE_ENV_ALLOWLIST


def run_ok(output_path: str, content: Any) -> subprocess.CompletedProcess[str]:
    """Return a subprocess.CompletedProcess that succeeds and writes *content* to output_path."""
    pathlib.Path(output_path).write_text(json.dumps(content))
    return subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")


def make_run_writing(content: Any) -> Callable[..., subprocess.CompletedProcess[str]]:
    """Return a mock_run side_effect that writes *content* to the --output path in the cmd."""

    def side_effect(cmd: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
        for i, arg in enumerate(cmd):
            if arg == "--output" and i + 1 < len(cmd):
                pathlib.Path(cmd[i + 1]).write_text(json.dumps(content))
                break
        return subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")

    return side_effect


class TestEvaluateAllCriteria:
    """Tests for evaluate_all_criteria IPC contract: dict, list, invalid shapes, failures."""

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_new_dict_shape(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """New object format: {verdicts: [...], llm_usage: {...}}."""
        mock_clone.return_value = str(tmp_path)
        output_content = {
            "verdicts": [
                {"index": 0, "met": True, "reasoning": "ok", "evidence": []},
                {"index": 1, "met": False, "reasoning": "no", "evidence": []},
            ],
            "llm_usage": {"cost_usd": 0.1, "prompt_tokens": 500},
        }

        mock_run.side_effect = make_run_writing(output_content)
        judge_input = make_batch_input(tmp_path, n=2)
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user="sandbox", trace_path=trace_path)

        assert len(verdicts) == 2
        assert verdicts[0].met is True
        assert verdicts[1].met is False
        assert usage.cost_usd == 0.1

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_nonzero_exit_returns_fail_all(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """Non-zero exit code from subprocess returns fail-all with empty usage."""
        mock_clone.return_value = str(tmp_path)
        mock_run.return_value = subprocess.CompletedProcess(args=[], returncode=1, stdout="", stderr="segfault")

        judge_input = make_batch_input(tmp_path, n=2)
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user="sandbox", trace_path=trace_path)

        assert len(verdicts) == 2
        assert all(v.met is None for v in verdicts)
        assert "exit 1" in verdicts[0].reasoning
        assert usage == LLMUsage()

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_timeout_returns_fail_all(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """Subprocess timeout returns fail-all with empty usage."""
        mock_clone.return_value = str(tmp_path)
        mock_run.side_effect = subprocess.TimeoutExpired(cmd="judge", timeout=300)

        judge_input = make_batch_input(tmp_path, n=2)
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user="sandbox", trace_path=trace_path)

        assert len(verdicts) == 2
        assert all(v.met is None for v in verdicts)
        assert "timed out" in verdicts[0].reasoning.lower()
        assert usage == LLMUsage()

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_invalid_json_in_output_file(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """Non-JSON content in output file returns fail-all."""
        mock_clone.return_value = str(tmp_path)

        def write_invalid(cmd: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
            for i, arg in enumerate(cmd):
                if arg == "--output" and i + 1 < len(cmd):
                    pathlib.Path(cmd[i + 1]).write_text("not valid json {{{")
                    break
            return subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")

        mock_run.side_effect = write_invalid

        judge_input = make_batch_input(tmp_path, n=1)
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user="sandbox", trace_path=trace_path)

        assert len(verdicts) == 1
        assert verdicts[0].met is None
        assert usage == LLMUsage()

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_empty_output_file(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """If the judge wrote nothing to the output file, return fail-all."""
        mock_clone.return_value = str(tmp_path)
        # mock_run does not write to the output file — it stays empty (pre-created by grader)
        mock_run.return_value = subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")

        judge_input = make_batch_input(tmp_path, n=2)
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user="sandbox", trace_path=trace_path)

        assert len(verdicts) == 2
        assert all(v.met is None for v in verdicts)
        assert usage == LLMUsage()


@pytest.fixture
def fake_judge(tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> pathlib.Path:
    """Create a fake ``gandalf-the-grader-judge`` on PATH.

    The script reads --input/--output args and writes a valid verdict to the
    output file.  When ``--batch`` is passed it writes the dict-with-verdicts
    format; otherwise a single-criterion verdict dict.
    """
    script = tmp_path / "bin" / "gandalf-the-grader-judge"
    script.parent.mkdir()
    script.write_text(
        """\
#!/usr/bin/env python3
import argparse, json, sys

parser = argparse.ArgumentParser()
parser.add_argument("--input", required=True)
parser.add_argument("--output", required=True)
parser.add_argument("--batch", action="store_true")
args = parser.parse_args()

inp = json.load(open(args.input))

if args.batch:
    verdicts = [
        {"index": i, "met": True, "reasoning": "ok", "evidence": []}
        for i in range(len(inp["criteria"]))
    ]
    result = {"verdicts": verdicts, "llm_usage": {"cost_usd": 0.01}}
else:
    result = {"verdict": {"met": True, "reasoning": "ok", "evidence": []}, "llm_usage": {"cost_usd": 0.01}}

with open(args.output, "w") as f:
    json.dump(result, f)
""",
    )
    script.chmod(0o755)
    monkeypatch.setenv("PATH", f"{script.parent}:{os.environ.get('PATH', '')}")
    return script


class TestSandboxUserNone:
    """When sandbox_user is None the judge runs as the ambient user (no sudo)."""

    @pytest.mark.usefixtures("fake_judge")
    def test_evaluate_criterion_no_sudo(self, tmp_path: pathlib.Path) -> None:
        workdir = tmp_path / "workspace"
        workdir.mkdir()
        (workdir / "hello.txt").write_text("hi")

        judge_input = JudgeInput(
            model="test-model",
            instructions="test",
            final_output="done",
            criterion="check something",
            workdir=str(workdir),
        )
        trace_path = str(tmp_path / "trace.txt")

        verdicts, _usage = run_judge(judge_input, sandbox_user=None, trace_path=trace_path)

        assert verdicts[0].met is True
        assert verdicts[0].reasoning == "ok"

    @pytest.mark.usefixtures("fake_judge")
    def test_evaluate_all_criteria_no_sudo(self, tmp_path: pathlib.Path) -> None:
        workdir = tmp_path / "workspace"
        workdir.mkdir()
        (workdir / "hello.txt").write_text("hi")

        judge_input = make_batch_input(tmp_path, n=2)
        judge_input = judge_input.model_copy(update={"workdir": str(workdir)})
        trace_path = str(tmp_path / "trace.txt")

        verdicts, usage = run_judge(judge_input, sandbox_user=None, trace_path=trace_path)

        assert len(verdicts) == 2
        assert all(v.met is True for v in verdicts)
        assert usage.cost_usd == 0.01


class TestSectionToleranceRuntime:
    def test_criterion_for_judge_appends_section_tolerance_for_non_gate(self) -> None:
        item = RubricItem(
            criterion="Revenue is $1,000",
            weight=1.0,
            section="Model Build",
            section_tolerance_pct=1.0,
        )

        criterion = criterion_for_judge(item, golden_check=False)

        assert criterion.startswith("Revenue is $1,000")
        assert "Section-level numeric tolerance" in criterion
        assert "+/-1%" in criterion
        assert "exact zero" in criterion

    def test_criterion_for_judge_does_not_append_for_gates_or_golden_check(self) -> None:
        gate = RubricItem(
            criterion="Workbook exists",
            weight=0.0,
            section="Model Build",
            gate=True,
            section_tolerance_pct=1.0,
        )
        normal = RubricItem(
            criterion="Revenue is $1,000",
            weight=1.0,
            section="Model Build",
            section_tolerance_pct=1.0,
        )

        assert criterion_for_judge(gate, golden_check=False) == "Workbook exists"
        assert criterion_for_judge(normal, golden_check=True) == "Revenue is $1,000"

    def test_effective_judge_guidance_adds_golden_check_instruction(self) -> None:
        guidance = effective_judge_guidance("Use openpyxl.", golden_check=True)

        assert guidance.startswith("Golden check mode is enabled.")
        assert "Section-level tolerance bands are disabled" in guidance
        assert guidance.endswith("Use openpyxl.")

    def test_effective_judge_guidance_adds_excel_backend_instruction(self) -> None:
        guidance = effective_judge_guidance(
            "Use openpyxl only for structure.",
            golden_check=False,
            excel_backend=ExcelBackendConfig(enabled=True),
        )

        assert "Microsoft Excel-backed workbook inspection" in guidance
        assert "LibreOffice" in guidance
        assert "include_formats=true" in guidance
        assert guidance.endswith("Use openpyxl only for structure.")

    def test_effective_judge_guidance_adds_libreoffice_backend_instruction(self) -> None:
        guidance = effective_judge_guidance(
            "Use MCP tools.",
            golden_check=False,
            excel_backend=ExcelBackendConfig(enabled=True, backend="libreoffice"),
        )

        assert "LibreOffice headless plus openpyxl" in guidance
        assert "include_formats=true" in guidance
        assert "Microsoft Excel-backed workbook inspection" not in guidance
        assert guidance.endswith("Use MCP tools.")

    @patch("gandalf.orchestrator.run_judge")
    def test_individual_mode_sends_augmented_criterion_but_preserves_result_text(
        self, mock_run_judge: Any, tmp_path: pathlib.Path
    ) -> None:
        config = make_config(workdir=str(tmp_path), output_dir=str(tmp_path / "output"), mode="individual")
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = [
            RubricItem(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                section_tolerance_pct=1.0,
            )
        ]
        mock_run_judge.return_value = ([Verdict(met=True, reasoning="ok")], LLMUsage())

        results, _usage = run_individual(config, rubric, "done", "test", "", None)

        judge_input = mock_run_judge.call_args[0][0]
        assert isinstance(judge_input, JudgeInput)
        assert "Section-level numeric tolerance" in judge_input.criterion
        assert results[0].criterion == "Revenue is $1,000"
        assert results[0].section_tolerance_pct == 1.0

    @patch("gandalf.orchestrator.run_judge")
    def test_batch_mode_sends_augmented_non_gate_only(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            reasoning_effort="xhigh",
            excel_backend=ExcelBackendConfig(enabled=True),
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = [
            RubricItem(
                criterion="Workbook exists",
                weight=0.0,
                section="Model Build",
                gate=True,
                section_tolerance_pct=1.0,
            ),
            RubricItem(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                section_tolerance_pct=1.0,
            ),
        ]
        mock_run_judge.return_value = (
            [Verdict(met=True, reasoning="ok"), Verdict(met=True, reasoning="ok")],
            LLMUsage(),
        )

        run_batch(config, rubric, "done", "test", "", None)

        judge_input = mock_run_judge.call_args[0][0]
        assert isinstance(judge_input, BatchJudgeInput)
        assert judge_input.reasoning_effort == "xhigh"
        assert judge_input.excel_backend.enabled is True
        assert judge_input.criteria[0] == "Workbook exists"
        assert "Section-level numeric tolerance" in judge_input.criteria[1]

    @patch("gandalf.orchestrator.run_judge")
    def test_batch_split_mode_sends_augmented_criteria(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = [
            RubricItem(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                section_tolerance_pct=1.0,
            ),
            RubricItem(criterion="Includes README", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()

        mock_run_judge.side_effect = _side_effect

        run_batch_concurrent(config, rubric, "done", "test", "", None)

        sent_criteria = [criterion for call in mock_run_judge.call_args_list for criterion in call[0][0].criteria]
        assert any("Section-level numeric tolerance" in criterion for criterion in sent_criteria)
        assert any(criterion == "Includes README" for criterion in sent_criteria)

    @patch("gandalf.orchestrator.run_judge")
    def test_golden_check_sends_unaugmented_criteria(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            golden_check=True,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = [
            RubricItem(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                section_tolerance_pct=1.0,
            )
        ]
        mock_run_judge.return_value = ([Verdict(met=True, reasoning="ok")], LLMUsage())

        run_batch(config, rubric, "done", "test", "Golden check mode is enabled.", None)

        judge_input = mock_run_judge.call_args[0][0]
        assert isinstance(judge_input, BatchJudgeInput)
        assert judge_input.criteria == ["Revenue is $1,000"]
        assert "Golden check mode is enabled." in judge_input.judge_guidance

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_cli_golden_check_overrides_config(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="batch",
            golden_check=False,
            auto_parallel=False,
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                section_tolerance_pct=1.0,
            )
        ]
        mock_run_batch.return_value = (
            [CriterionResult(criterion="Revenue is $1,000", weight=1.0, met=True, reasoning="ok")],
            LLMUsage(),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml", "--golden-check"]):
            main()

        call_args = mock_run_batch.call_args[0]
        assert call_args[0].golden_check is True
        assert call_args[0].workbook_digest.enabled is True
        assert call_args[0].rate_limit_discovery is True
        assert "Golden check mode is enabled." in call_args[4]


class TestWorkbookRepairCheck:
    def test_default_repair_check_workbook_prefers_submitted_output(self, tmp_path: pathlib.Path) -> None:
        assert default_repair_check_workbook(str(tmp_path)) is None
        (tmp_path / "submitted_output.xlsx").write_bytes(b"workbook")

        assert default_repair_check_workbook(str(tmp_path)) == "submitted_output.xlsx"

    def test_run_workbook_repair_check_skips_when_no_default_workbook(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))

        result = run_workbook_repair_check(config)

        assert result is not None
        assert result.checked is False
        assert result.skipped_reason == "submitted_output.xlsx not found"

    def test_run_workbook_repair_check_uses_excel_service(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        (tmp_path / "submitted_output.xlsx").write_bytes(b"workbook")
        config = make_config(workdir=str(tmp_path))

        class FakeService:
            def workbook_repair_check(self, path: str) -> dict[str, Any]:
                return {
                    "path": path,
                    "backend": "mac_excel",
                    "checked": True,
                    "opened": True,
                    "repair_dialog_detected": True,
                }

        monkeypatch.setattr("gandalf.orchestrator.create_service", lambda _config: FakeService())

        result = run_workbook_repair_check(config)

        assert result is not None
        assert result.checked is True
        assert result.workbook_path == "submitted_output.xlsx"
        assert result.backend == "mac_excel"
        assert result.repair_dialog_detected is True

    @staticmethod
    def _capturing_create_service(captured: dict[str, Any]) -> Any:
        class FakeService:
            def __init__(self, workdir: pathlib.Path) -> None:
                self.workdir = workdir

            def workbook_repair_check(self, path: str) -> dict[str, Any]:
                copied_path = self.workdir / path
                captured["path"] = path
                captured["workdir"] = self.workdir
                captured["copied_bytes"] = copied_path.read_bytes()
                return {
                    "path": path,
                    "backend": "mac_excel",
                    "checked": True,
                    "opened": True,
                    "repair_dialog_detected": False,
                }

        def fake_create_service(service_config: Any) -> FakeService:
            return FakeService(pathlib.Path(service_config.workdir))

        return fake_create_service

    def test_run_workbook_repair_check_stages_under_excel_access(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        (tmp_path / "submitted_output.xlsx").write_bytes(b"workbook")
        config = make_config(workdir=str(tmp_path))
        excel_access = tmp_path / "gandalf_runs" / "excel_access"
        captured: dict[str, Any] = {}
        monkeypatch.setattr("gandalf.orchestrator.excel_access_root", lambda: excel_access)
        monkeypatch.setattr("gandalf.orchestrator.create_service", self._capturing_create_service(captured))

        result = run_workbook_repair_check(config)

        assert result is not None
        assert result.workbook_path == "submitted_output.xlsx"
        assert captured["copied_bytes"] == b"workbook"
        # Mac Excel always opens from the stable excel_access root (grant-once),
        # under a per-run subdir that is cleaned up afterwards.
        assert captured["workdir"] != tmp_path
        assert captured["workdir"].is_relative_to(excel_access / "repair")
        assert not captured["workdir"].exists()

    def test_run_workbook_repair_check_keeps_windows_vm_original_path(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        (tmp_path / "submitted_output.xlsx").write_bytes(b"workbook")
        config = make_config(
            workdir=str(tmp_path),
            excel_backend={"backend": "windows_vm", "windows_url": "http://excel-vm.local"},
        )
        captured: dict[str, Any] = {}

        class FakeService:
            def workbook_repair_check(self, path: str) -> dict[str, Any]:
                captured["path"] = path
                return {
                    "path": path,
                    "backend": "windows_vm",
                    "checked": True,
                    "opened": True,
                    "repair_dialog_detected": False,
                }

        def fake_create_service(service_config: Any) -> FakeService:
            captured["workdir"] = pathlib.Path(service_config.workdir)
            return FakeService()

        monkeypatch.setattr("gandalf.orchestrator.create_service", fake_create_service)

        result = run_workbook_repair_check(config)

        assert result is not None
        assert result.workbook_path == "submitted_output.xlsx"
        assert captured == {"path": "submitted_output.xlsx", "workdir": tmp_path}

    def test_run_workbook_repair_check_respects_disabled_config(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path), workbook_repair_check={"enabled": False})

        result = run_workbook_repair_check(config)

        assert result is not None
        assert result.enabled is False
        assert result.checked is False
        assert result.skipped_reason == "disabled"


class TestExcelPreflightCli:
    """Tests for the CLI-only Excel permission preflight."""

    def test_run_excel_preflight_skips_disabled_backend(
        self,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        config = make_config(
            workdir=str(tmp_path),
            excel_backend={"enabled": False},
        )

        assert run_excel_preflight(config) == 0
        assert "Excel preflight skipped" in capsys.readouterr().out

    def test_run_excel_preflight_opens_configured_workbook(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        (tmp_path / "model.xlsx").write_bytes(b"workbook")
        config = make_config(
            workdir=str(tmp_path),
            workbook_repair_check={"workbook_path": "model.xlsx"},
        )
        captured: dict[str, Any] = {}

        class FakeService:
            def workbook_summary(self, path: str, *, recalculate: bool) -> dict[str, Any]:
                captured["path"] = path
                captured["recalculate"] = recalculate
                return {
                    "path": path,
                    "backend": "mac_excel",
                    "sheets": [{"name": "Sheet1"}],
                }

        monkeypatch.setattr("gandalf.orchestrator.create_service", lambda _config: FakeService())

        assert run_excel_preflight(config) == 0
        assert captured == {"path": "model.xlsx", "recalculate": False}
        assert "Excel preflight passed for model.xlsx" in capsys.readouterr().out

    def test_run_excel_preflight_errors_without_workbook(
        self,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        config = make_config(workdir=str(tmp_path))

        assert run_excel_preflight(config) == 1
        assert "submitted_output.xlsx not found" in capsys.readouterr().err

    @patch("gandalf.orchestrator.resolve_instructions")
    @patch("gandalf.orchestrator.load_trajectory_final_output")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.create_service")
    def test_main_excel_preflight_exits_before_grading(
        self,
        mock_create_service: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,
        mock_instructions: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        (tmp_path / "submitted_output.xlsx").write_bytes(b"workbook")
        mock_config.return_value = make_config(workdir=str(tmp_path))

        class FakeService:
            def workbook_summary(self, path: str, *, recalculate: bool) -> dict[str, Any]:
                return {"path": path, "backend": "mac_excel", "sheets": [{"name": "Sheet1"}]}

        mock_create_service.return_value = FakeService()

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--excel-preflight"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 0
        mock_rubric.assert_not_called()
        mock_trajectory.assert_not_called()
        mock_instructions.assert_not_called()

    @pytest.mark.parametrize(
        "extra_args",
        [
            ["--gates-only"],
            ["--include-section", "Section A"],
            ["--exclude-section", "Section A"],
        ],
    )
    @patch("gandalf.orchestrator.load_config")
    def test_main_excel_preflight_rejects_rubric_modes(
        self,
        mock_config: Any,
        extra_args: list[str],
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--excel-preflight", *extra_args]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 2
        assert "--excel-preflight cannot be combined with" in capsys.readouterr().err
        mock_config.assert_not_called()


class TestExcelBackendCliOverride:
    """Tests for per-run Excel backend selection."""

    def test_libreoffice_override_enables_libreoffice_backend(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))

        overridden = apply_excel_backend_cli_override(config, backend="libreoffice")

        assert overridden.excel_backend.enabled is True
        assert overridden.excel_backend.backend == "libreoffice"

    def test_disabled_override_disables_excel_backend(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))

        overridden = apply_excel_backend_cli_override(config, backend="disabled")

        assert overridden.excel_backend.enabled is False
        assert overridden.excel_backend.backend == "mac_excel"

    def test_windows_vm_override_uses_cli_url_and_token(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))

        overridden = apply_excel_backend_cli_override(
            config,
            backend="windows-vm",
            windows_url="http://excel-vm.local:8765/",
            auth_token=" secret ",
        )

        assert overridden.excel_backend.enabled is True
        assert overridden.excel_backend.backend == "windows_vm"
        assert overridden.excel_backend.windows_url == "http://excel-vm.local:8765"
        assert overridden.excel_backend.auth_token == "secret"

    def test_windows_vm_override_requires_url(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))

        with pytest.raises(ValueError, match="windows_url"):
            apply_excel_backend_cli_override(config, backend="windows-vm")

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_libreoffice_override_enables_libreoffice_guidance_and_tools(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        (tmp_path / "submitted_output.xlsx").write_bytes(b"placeholder")
        mock_config.return_value = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            auto_parallel=False,
        )
        mock_rubric.return_value = [RubricItem(criterion="a", weight=1.0)]

        def _run(
            config: GraderConfig,
            rubric: list[RubricItem],
            _final_output: str,
            _instructions: str,
            judge_guidance: str,
            _judge_prompt: str | None,
            **_kwargs: Any,
        ) -> tuple[list[CriterionResult], LLMUsage]:
            assert config.excel_backend.enabled is True
            assert config.excel_backend.backend == "libreoffice"
            assert "LibreOffice headless plus openpyxl" in judge_guidance
            assert "Microsoft Excel-backed workbook inspection" not in judge_guidance
            return [
                CriterionResult(
                    criterion=rubric[0].criterion,
                    weight=rubric[0].weight,
                    met=True,
                    reasoning="ok",
                )
            ], LLMUsage()

        mock_run_batch.side_effect = _run

        with patch("sys.argv", ["prog", "--config", "dummy.toml", "--excel-backend", "libreoffice"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["workbook_repair_check"]["skipped_reason"] == "repair dialogue detection requires Microsoft Excel"
        assert info["workbook_repair_check"]["backend"] == "libreoffice"

    @patch("gandalf.orchestrator.resolve_instructions")
    @patch("gandalf.orchestrator.load_trajectory_final_output")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_windows_vm_override_errors_before_grading_without_url(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,
        mock_instructions: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = make_config(workdir=str(tmp_path), output_dir=str(tmp_path / "output"))

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--excel-backend", "windows-vm"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 2
        assert "windows_url" in capsys.readouterr().err
        mock_rubric.assert_not_called()
        mock_trajectory.assert_not_called()
        mock_instructions.assert_not_called()


class TestCollectExcelMetrics:
    def test_returns_none_when_no_sidecars(self, tmp_path: pathlib.Path) -> None:
        assert collect_excel_metrics(str(tmp_path)) is None

    def test_merges_sidecars_into_one_summary(self, tmp_path: pathlib.Path) -> None:
        (tmp_path / "judge_trace_0_excel_metrics.json").write_text(
            json.dumps({"call_count": 2, "operation_counts": {"workbook_summary": 2}})
        )
        (tmp_path / "judge_trace_1_excel_metrics.json").write_text(
            json.dumps({"call_count": 3, "operation_counts": {"find_cells": 3}})
        )

        merged = collect_excel_metrics(str(tmp_path))

        assert merged is not None
        assert merged["call_count"] == 5
        assert merged["operation_counts"] == {"workbook_summary": 2, "find_cells": 3}

    def test_ignores_malformed_sidecar(self, tmp_path: pathlib.Path) -> None:
        (tmp_path / "a_excel_metrics.json").write_text("{not json")
        (tmp_path / "b_excel_metrics.json").write_text(json.dumps({"call_count": 1}))

        merged = collect_excel_metrics(str(tmp_path))

        assert merged is not None
        assert merged["call_count"] == 1


class TestLibreOfficeCacheEnvironment:
    def test_creates_and_cleans_cache_dir_for_libreoffice(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.delenv("GANDALF_EXCEL_CACHE_DIR", raising=False)
        config = make_config(
            workdir=str(tmp_path),
            excel_backend={"enabled": True, "backend": "libreoffice"},
        )

        with libreoffice_cache_environment(config):
            cache_dir = pathlib.Path(os.environ["GANDALF_EXCEL_CACHE_DIR"])
            assert cache_dir.is_dir()
            assert cache_dir.parent == pathlib.Path(tempfile.gettempdir())

        assert "GANDALF_EXCEL_CACHE_DIR" not in os.environ
        assert not cache_dir.exists()

    def test_restores_previous_cache_dir(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.setenv("GANDALF_EXCEL_CACHE_DIR", "/tmp/existing-lo-cache")
        config = make_config(
            workdir=str(tmp_path),
            excel_backend={"enabled": True, "backend": "libreoffice"},
        )

        with libreoffice_cache_environment(config):
            assert os.environ["GANDALF_EXCEL_CACHE_DIR"] != "/tmp/existing-lo-cache"

        assert os.environ["GANDALF_EXCEL_CACHE_DIR"] == "/tmp/existing-lo-cache"

    def test_non_libreoffice_does_not_set_cache_dir(
        self,
        tmp_path: pathlib.Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        monkeypatch.delenv("GANDALF_EXCEL_CACHE_DIR", raising=False)
        config = make_config(
            workdir=str(tmp_path),
            excel_backend={"enabled": True, "backend": "mac_excel"},
        )

        with libreoffice_cache_environment(config):
            assert "GANDALF_EXCEL_CACHE_DIR" not in os.environ


class TestWorkbookDigest:
    def test_disabled_by_default_returns_none(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path))
        assert workbook_digest_text(config) is None

    def test_enabled_returns_digest_for_default_workbook(self, tmp_path: pathlib.Path) -> None:
        wb = Workbook()
        wb.active.title = "Summary"
        wb.active["A1"] = "Revenue"
        wb.active["B1"] = "=A1"
        wb.save(tmp_path / "submitted_output.xlsx")
        config = make_config(workdir=str(tmp_path), workbook_digest={"enabled": True})

        digest = workbook_digest_text(config)

        assert digest is not None
        assert "Workbook digest for submitted_output.xlsx" in digest
        assert '"Summary" (visible)' in digest

    def test_enabled_but_missing_workbook_returns_none(self, tmp_path: pathlib.Path) -> None:
        config = make_config(workdir=str(tmp_path), workbook_digest={"enabled": True})
        assert workbook_digest_text(config) is None

    def test_rejects_relative_path_outside_workdir(self, tmp_path: pathlib.Path) -> None:
        workdir = tmp_path / "workdir"
        workdir.mkdir()
        outside = tmp_path / "reference.xlsx"
        Workbook().save(outside)
        config = make_config(
            workdir=str(workdir),
            workbook_digest={"enabled": True, "workbook_path": "../reference.xlsx"},
        )

        with pytest.raises(ValueError, match="outside the allowed workdir"):
            workbook_digest_text(config)

    def test_rejects_absolute_path_outside_workdir(self, tmp_path: pathlib.Path) -> None:
        workdir = tmp_path / "workdir"
        workdir.mkdir()
        outside = tmp_path / "reference.xlsx"
        Workbook().save(outside)
        config = make_config(
            workdir=str(workdir),
            workbook_digest={"enabled": True, "workbook_path": str(outside)},
        )

        with pytest.raises(ValueError, match="outside the allowed workdir"):
            workbook_digest_text(config)

    def test_effective_judge_guidance_includes_digest(self) -> None:
        guidance = effective_judge_guidance(
            "extra guidance",
            golden_check=False,
            workbook_digest="DIGEST-BLOCK",
        )
        assert "DIGEST-BLOCK" in guidance
        assert "extra guidance" in guidance


class TestGoldenCheckDefaults:
    def test_golden_check_enables_digest_by_default(self) -> None:
        config = make_config(golden_check=True)

        adjusted = apply_golden_check_defaults(config)

        assert adjusted.workbook_digest.enabled is True

    def test_golden_check_respects_explicit_digest_override(self) -> None:
        config = make_config(
            golden_check=True,
            workbook_digest={"enabled": False},
        )

        adjusted = apply_golden_check_defaults(config)

        assert adjusted.workbook_digest.enabled is False

    def test_non_golden_check_leaves_digest_default_disabled(self) -> None:
        config = make_config()

        adjusted = apply_golden_check_defaults(config)

        assert adjusted.workbook_digest.enabled is False


class TestExcelWorkspaceWarning:
    def test_warns_when_mac_excel_and_output_dir_protected(
        self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        config = make_config(output_dir=str(tmp_path / "out"), workdir=str(tmp_path / "wd"))
        monkeypatch.setattr(
            "gandalf.orchestrator.is_protected_path",
            lambda path: path == config.output_dir,
        )
        warn_if_excel_workspace_protected(config)
        err = capsys.readouterr().err
        assert "WARNING" in err
        assert "output_dir" in err
        assert "workdir" not in err

    def test_no_warning_when_paths_unprotected(
        self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        config = make_config(output_dir=str(tmp_path / "out"), workdir=str(tmp_path / "wd"))
        monkeypatch.setattr("gandalf.orchestrator.is_protected_path", lambda _path: False)
        warn_if_excel_workspace_protected(config)
        assert capsys.readouterr().err == ""

    def test_no_warning_when_backend_not_mac_excel(
        self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        config = make_config(
            output_dir=str(tmp_path / "out"),
            workdir=str(tmp_path / "wd"),
            excel_backend={"backend": "windows_vm", "windows_url": "http://excel-vm.local"},
        )
        monkeypatch.setattr("gandalf.orchestrator.is_protected_path", lambda _path: True)
        warn_if_excel_workspace_protected(config)
        assert capsys.readouterr().err == ""


class TestScoring:
    """Tests for write_info scoring: raw_score, reward, and bounds.

    Each test asserts both raw_score and reward together so the full
    scoring pipeline is verified in one place per scenario.
    """

    def info(self, results: list[CriterionResult], tmp_path: pathlib.Path) -> dict[str, Any]:
        """Run write_info and return parsed info.json."""
        config = make_config(output_dir=str(tmp_path))
        errored = sum(1 for r in results if r.met is None)
        write_info(config, results, LLMUsage(), errored)
        with open(tmp_path / "info.json") as f:
            result: dict[str, Any] = json.load(f)
            return result

    # -- core scenarios (raw_score + reward together) --

    def test_all_positive_all_met(self, tmp_path: pathlib.Path) -> None:
        """weights=[2,3], met=[T,T] → raw=5, reward=1.0, min=0, max=5."""
        info = self.info([cr(weight=2.0, met=True), cr(weight=3.0, met=True)], tmp_path)
        assert info["raw_score"] == 5.0
        assert info["reward"] == 1.0
        assert info["minimum_score"] == 0.0
        assert info["maximum_score"] == 5.0

    def test_all_positive_partial_met(self, tmp_path: pathlib.Path) -> None:
        """weights=[2,3], met=[T,F] → raw=2, reward=0.4."""
        info = self.info([cr(weight=2.0, met=True), cr(weight=3.0, met=False)], tmp_path)
        assert info["raw_score"] == 2.0
        assert info["reward"] == 0.4

    def test_all_positive_none_met(self, tmp_path: pathlib.Path) -> None:
        """weights=[2,3], met=[F,F] → raw=0, reward=0.0."""
        info = self.info([cr(weight=2.0, met=False), cr(weight=3.0, met=False)], tmp_path)
        assert info["raw_score"] == 0.0
        assert info["reward"] == 0.0

    def test_mixed_negative_penalty_applied(self, tmp_path: pathlib.Path) -> None:
        """weights=[3,-1], met=[T,T] → raw=2, reward=2/3."""
        info = self.info([cr(weight=3.0, met=True), cr(weight=-1.0, met=True)], tmp_path)
        assert info["raw_score"] == 2.0
        assert info["reward"] == 0.6667

    def test_mixed_negative_drives_below_zero_clipped(self, tmp_path: pathlib.Path) -> None:
        """weights=[1,-3], met=[F,T] → raw=-3, reward=0.0 (clip lower bound)."""
        info = self.info([cr(weight=1.0, met=False), cr(weight=-3.0, met=True)], tmp_path)
        assert info["raw_score"] == -3.0
        assert info["reward"] == 0.0

    def test_negative_not_met_no_penalty(self, tmp_path: pathlib.Path) -> None:
        """weights=[3,-1], met=[T,F] → raw=3, reward=1.0."""
        info = self.info([cr(weight=3.0, met=True), cr(weight=-1.0, met=False)], tmp_path)
        assert info["raw_score"] == 3.0
        assert info["reward"] == 1.0

    def test_all_negative_denominator_zero(self, tmp_path: pathlib.Path) -> None:
        """weights=[-2,-3], met=[T,T] → raw=-5, reward=0.0 (no divide-by-zero)."""
        info = self.info([cr(weight=-2.0, met=True), cr(weight=-3.0, met=True)], tmp_path)
        assert info["raw_score"] == -5.0
        assert info["reward"] == 0.0

    def test_empty_rubric(self, tmp_path: pathlib.Path) -> None:
        """No criteria → raw=0, reward=0."""
        info = self.info([], tmp_path)
        assert info["raw_score"] == 0.0
        assert info["reward"] == 0.0

    def test_errored_positive_criterion(self, tmp_path: pathlib.Path) -> None:
        """weights=[3,2], met=[T,None] → raw=3, reward=3/5=0.6."""
        info = self.info([cr(weight=3.0, met=True), cr(weight=2.0, met=None)], tmp_path)
        assert info["raw_score"] == 3.0
        assert info["reward"] == 0.6

    def test_errored_negative_criterion(self, tmp_path: pathlib.Path) -> None:
        """weights=[3,-2], met=[T,None] → raw=3, reward=3/3=1.0."""
        info = self.info([cr(weight=3.0, met=True), cr(weight=-2.0, met=None)], tmp_path)
        assert info["raw_score"] == 3.0
        assert info["reward"] == 1.0

    # -- info.json shape --

    def test_info_json_contains_reward_and_raw_score(self, tmp_path: pathlib.Path) -> None:
        """info.json must contain both reward and raw_score fields."""
        info = self.info([cr(weight=2.0, met=True), cr(weight=3.0, met=False)], tmp_path)
        assert "reward" in info
        assert "raw_score" in info
        assert isinstance(info["reward"], float)
        assert isinstance(info["raw_score"], (int, float))

    def test_info_json_contains_minimum_and_maximum_score(self, tmp_path: pathlib.Path) -> None:
        info = self.info([cr(weight=10.0, met=True), cr(weight=5.0, met=False), cr(weight=-3.0, met=True)], tmp_path)
        assert info["minimum_score"] == -3.0
        assert info["maximum_score"] == 15.0

    def test_info_json_includes_excel_metrics_when_present(self, tmp_path: pathlib.Path) -> None:
        config = make_config(output_dir=str(tmp_path))

        write_info(
            config,
            [cr(weight=1.0, met=True)],
            LLMUsage(),
            errored_criterion_count=0,
            excel_metrics={
                "call_count": 2,
                "operation_counts": {"workbook_summary": 1, "find_cells": 1},
                "excel_service_cache_hits": 1,
            },
        )

        info = json.loads((tmp_path / "info.json").read_text())
        assert info["excel_metrics"]["call_count"] == 2
        assert info["excel_metrics"]["operation_counts"]["find_cells"] == 1

    def test_repair_check_true_halves_reward_only(self, tmp_path: pathlib.Path) -> None:
        config = make_config(output_dir=str(tmp_path))
        repair_check = WorkbookRepairCheckResult(
            checked=True,
            workbook_path="submitted_output.xlsx",
            backend="mac_excel",
            opened=True,
            repair_dialog_detected=True,
        )

        reward, raw_score = write_info(
            config,
            [cr(weight=2.0, met=True), cr(weight=3.0, met=True)],
            LLMUsage(),
            errored_criterion_count=0,
            workbook_repair_check=repair_check,
        )

        info = json.loads((tmp_path / "info.json").read_text())
        assert raw_score == 5.0
        assert reward == 0.5
        assert info["raw_score"] == 5.0
        assert info["unpenalized_reward"] == 1.0
        assert info["reward"] == 0.5
        assert info["repair_penalty_applied"] is True
        assert info["repair_penalty_multiplier"] == 0.5
        assert info["workbook_repair_check"]["repair_dialog_detected"] is True

    def test_repair_check_false_records_no_penalty(self, tmp_path: pathlib.Path) -> None:
        config = make_config(output_dir=str(tmp_path))
        repair_check = WorkbookRepairCheckResult(
            checked=True,
            workbook_path="submitted_output.xlsx",
            backend="mac_excel",
            opened=True,
            repair_dialog_detected=False,
        )

        reward, _raw_score = write_info(
            config,
            [cr(weight=2.0, met=True), cr(weight=3.0, met=True)],
            LLMUsage(),
            errored_criterion_count=0,
            workbook_repair_check=repair_check,
        )

        info = json.loads((tmp_path / "info.json").read_text())
        assert reward == 1.0
        assert info["reward"] == 1.0
        assert info["repair_penalty_applied"] is False
        assert "unpenalized_reward" not in info
        assert "repair_penalty_multiplier" not in info
        assert info["workbook_repair_check"]["repair_dialog_detected"] is False

    def test_legacy_flat_rubric_has_empty_section_results(self, tmp_path: pathlib.Path) -> None:
        info = self.info([cr(weight=2.0, met=True), cr(weight=3.0, met=False)], tmp_path)
        assert info["raw_score"] == 2.0
        assert info["section_results"] == []
        assert info["criterion_results"][0]["section"] is None
        assert info["criterion_results"][0]["gate"] is False
        assert info["criterion_results"][0]["score_contribution"] == 2.0
        assert "section_tolerance_pct" not in info["criterion_results"][0]

    def test_info_json_includes_section_tolerance_in_normal_mode(self, tmp_path: pathlib.Path) -> None:
        results = [
            CriterionResult(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                met=True,
                reasoning="ok",
                section_tolerance_pct=1.0,
            )
        ]
        config = make_config(output_dir=str(tmp_path))

        write_info(config, results, LLMUsage(), errored_criterion_count=0)

        info = json.loads((tmp_path / "info.json").read_text())
        assert info["criterion_results"][0]["section_tolerance_pct"] == 1.0
        assert info["section_results"][0]["section_tolerance_pct"] == 1.0

    def test_info_json_omits_section_tolerance_in_golden_check(self, tmp_path: pathlib.Path) -> None:
        results = [
            CriterionResult(
                criterion="Revenue is $1,000",
                weight=1.0,
                section="Model Build",
                met=True,
                reasoning="ok",
                section_tolerance_pct=1.0,
            )
        ]
        config = make_config(output_dir=str(tmp_path), golden_check=True)

        write_info(config, results, LLMUsage(), errored_criterion_count=0)

        info = json.loads((tmp_path / "info.json").read_text())
        assert "section_tolerance_pct" not in info["criterion_results"][0]
        assert "section_tolerance_pct" not in info["section_results"][0]

    def test_gated_section_scores_when_all_gates_pass(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=True, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
                cr(weight=2.0, met=False, section="Model Build"),
            ],
            tmp_path,
        )

        assert info["raw_score"] == 5.0
        assert info["reward"] == 0.7143
        assert info["criterion_results"][1]["score_contribution"] == 5.0
        assert info["section_results"] == [
            {
                "section": "Model Build",
                "score": 5.0,
                "minimum_score": 0.0,
                "maximum_score": 7.0,
                "gate_count": 1,
                "passed_gate_indices": [0],
                "failed_gate_indices": [],
                "section_gate": None,
                "section_gate_met": None,
                "section_gate_reasoning": None,
                "section_gate_evidence": [],
                "section_gate_count": 0,
                "section_gate_results": [],
                "passed_section_gate_indices": [],
                "failed_section_gate_indices": [],
                "errored_section_gate_indices": [],
                "skipped_section_gate_indices": [],
                "section_gates_met": None,
                "gated_out": False,
            }
        ]

    def test_failed_gate_zeroes_entire_section_including_penalties(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=False, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
                cr(weight=-2.0, met=True, section="Model Build"),
            ],
            tmp_path,
        )

        assert info["raw_score"] == 0.0
        assert info["reward"] == 0.0
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["criterion_results"][2]["score_contribution"] == 0.0
        assert info["section_results"][0]["score"] == 0.0
        assert info["section_results"][0]["failed_gate_indices"] == [0]
        assert info["section_results"][0]["gated_out"] is True

    def test_failed_gate_does_not_suppress_other_section_penalties(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=False, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
                cr(weight=-2.0, met=True, section="Model Build"),
                cr(weight=-3.0, met=True, section="Penalties"),
            ],
            tmp_path,
        )

        assert info["raw_score"] == -3.0
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["criterion_results"][2]["score_contribution"] == 0.0
        assert info["criterion_results"][3]["score_contribution"] == -3.0
        assert info["section_results"][0]["section"] == "Model Build"
        assert info["section_results"][0]["score"] == 0.0
        assert info["section_results"][1]["section"] == "Penalties"
        assert info["section_results"][1]["score"] == -3.0

    def test_multiple_gates_require_all_to_pass(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=True, section="Model Build", gate=True),
                cr(weight=0.0, met=False, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
            ],
            tmp_path,
        )

        assert info["raw_score"] == 0.0
        assert info["criterion_results"][2]["score_contribution"] == 0.0
        assert info["section_results"][0]["gate_count"] == 2
        assert info["section_results"][0]["passed_gate_indices"] == [0]
        assert info["section_results"][0]["failed_gate_indices"] == [1]
        assert info["section_results"][0]["gated_out"] is True

    def test_errored_gate_gates_out_section(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=None, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
            ],
            tmp_path,
        )

        assert info["raw_score"] == 0.0
        assert info["errored_criterion_count"] == 1
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["section_results"][0]["failed_gate_indices"] == [0]
        assert info["section_results"][0]["gated_out"] is True

    def test_sections_without_gates_score_normally(self, tmp_path: pathlib.Path) -> None:
        info = self.info(
            [
                cr(weight=0.0, met=False, section="Model Build", gate=True),
                cr(weight=5.0, met=True, section="Model Build"),
                cr(weight=3.0, met=True, section="Formatting"),
                cr(weight=2.0, met=True),
            ],
            tmp_path,
        )

        assert info["raw_score"] == 5.0
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["criterion_results"][2]["score_contribution"] == 3.0
        assert info["criterion_results"][3]["score_contribution"] == 2.0
        assert info["section_results"][0]["section"] == "Model Build"
        assert info["section_results"][0]["gated_out"] is True
        assert info["section_results"][1]["section"] == "Formatting"
        assert info["section_results"][1]["gate_count"] == 0
        assert info["section_results"][1]["gated_out"] is False

    def test_failed_section_gate_zeroes_entire_section(self, tmp_path: pathlib.Path) -> None:
        results = [
            CriterionResult(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                met=None,
                reasoning="Skipped because section gate failed.",
                skipped=True,
                skip_reason="Skipped because section gate failed.",
            ),
            CriterionResult(
                criterion="Used hardcoded values",
                weight=-2.0,
                section="Model Build",
                met=None,
                reasoning="Skipped because section gate failed.",
                skipped=True,
                skip_reason="Skipped because section gate failed.",
            ),
        ]
        section_gate_results = {
            "Model Build": [
                SectionGateResult(
                    section="Model Build",
                    index=0,
                    criterion="The submitted workbook exists.",
                    met=False,
                    reasoning="Workbook missing.",
                    evidence=["No workbook found"],
                )
            ]
        }
        config = make_config(output_dir=str(tmp_path))
        write_info(config, results, LLMUsage(), errored_criterion_count=0, section_gate_results=section_gate_results)

        info = json.loads((tmp_path / "info.json").read_text())
        assert info["raw_score"] == 0.0
        assert info["evaluated_criteria_pct"] == 0.0
        assert info["criterion_results"][0]["skipped"] is True
        assert info["criterion_results"][1]["skipped"] is True
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["section_results"][0]["score"] == 0.0
        assert info["section_results"][0]["section_gate"] == "The submitted workbook exists."
        assert info["section_results"][0]["section_gate_met"] is False
        assert info["section_results"][0]["section_gate_count"] == 1
        assert info["section_results"][0]["section_gate_results"][0]["criterion"] == "The submitted workbook exists."
        assert info["section_results"][0]["failed_section_gate_indices"] == [0]
        assert info["section_results"][0]["section_gates_met"] is False
        assert info["section_results"][0]["gated_out"] is True


class TestOutputFilePermissions:
    """Ensure the judge output file is pre-created with world-writable permissions.

    Regression: the old code used tempfile.mktemp() which does NOT create the
    file, requiring sandbox_user to create it in /tmp.  On systems where /tmp
    is not world-writable, this caused a PermissionError.  The fix pre-creates
    the file and chmods it 0o666 so sandbox_user only needs to *write* to an
    existing file, not *create* one in a restricted directory.
    """

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_output_file_exists_before_subprocess(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """Output file must be pre-created so sandbox_user can write it without /tmp access."""
        mock_clone.return_value = str(tmp_path)
        captured_cmd: dict[str, Any] = {}

        def capture(cmd: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
            output_path = cmd[cmd.index("--output") + 1]
            captured_cmd["output_path"] = output_path
            captured_cmd["existed_before_run"] = pathlib.Path(output_path).exists()
            # Simulate sandbox_user writing to the pre-created file
            pathlib.Path(output_path).write_text(
                json.dumps(
                    {
                        "verdicts": [{"met": True, "reasoning": "ok", "evidence": []}],
                        "llm_usage": {
                            "cost_usd": 0,
                            "prompt_tokens": 0,
                            "completion_tokens": 0,
                            "cache_read_tokens": 0,
                        },
                    }
                )
            )
            return subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")

        mock_run.side_effect = capture
        judge_input = make_batch_input(tmp_path, n=1)

        run_judge(judge_input, sandbox_user="sandbox", trace_path=str(tmp_path / "trace.txt"))

        assert "output_path" in captured_cmd, "subprocess was not called"
        assert captured_cmd.get("existed_before_run"), (
            "Output file was NOT pre-created before subprocess.run — "
            "sandbox_user would need to create it in /tmp (may not be world-writable)"
        )

    @patch("gandalf.orchestrator.clone_workspace")
    @patch("gandalf.orchestrator.subprocess.run")
    def test_output_file_is_world_writable(self, mock_run: Any, mock_clone: Any, tmp_path: pathlib.Path) -> None:
        """Pre-created output file must have world-write so sandbox_user can overwrite it.

        This test fails on the pre-fix code (tempfile.mktemp → file never created)
        and passes with the fix (NamedTemporaryFile + chmod 0o666).
        """
        mock_clone.return_value = str(tmp_path)
        captured: dict[str, Any] = {}

        def capture_and_check_permissions(cmd: list[str], **_kwargs: Any) -> subprocess.CompletedProcess[str]:
            output_path = cmd[cmd.index("--output") + 1]
            captured["path"] = output_path
            captured["exists"] = pathlib.Path(output_path).exists()
            if captured["exists"]:
                captured["mode"] = os.stat(output_path).st_mode
            pathlib.Path(output_path).write_text(
                json.dumps(
                    {
                        "verdicts": [{"met": True, "reasoning": "ok", "evidence": []}],
                        "llm_usage": {
                            "cost_usd": 0,
                            "prompt_tokens": 0,
                            "completion_tokens": 0,
                            "cache_read_tokens": 0,
                        },
                    }
                )
            )
            return subprocess.CompletedProcess(args=[], returncode=0, stdout="", stderr="")

        mock_run.side_effect = capture_and_check_permissions
        judge_input = make_batch_input(tmp_path, n=1)
        run_judge(judge_input, sandbox_user="sandbox", trace_path=str(tmp_path / "trace.txt"))

        assert captured.get("exists"), (
            "Output file was NOT pre-created before subprocess.run — "
            "sandbox_user would need to create it in /tmp (may not be world-writable)"
        )
        mode = captured.get("mode", 0)
        assert mode & 0o002, (
            f"Output file missing world-write bit (mode={oct(mode)}) — "
            "sandbox_user cannot write to it without /tmp create access"
        )


class TestRetryLogic:
    """Tests for retry and hard-fail logic in main()."""

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_sequential_retry_resolves_errored_criterion(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Sequential retry resolves an errored criterion on the second attempt."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="individual",
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
        ]

        # First call: c1 passes, c2 errors. Retry: c2 passes.
        mock_eval.side_effect = [
            ([Verdict(met=True, reasoning="ok", evidence=["e1"])], LLMUsage()),
            ([Verdict(met=None, reasoning="timeout")], LLMUsage()),
            # retry for c2
            ([Verdict(met=True, reasoning="ok on retry", evidence=["e2"])], LLMUsage()),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["criterion_results"][0]["met"] is True
        assert info["criterion_results"][1]["met"] is True
        assert info["errored_criterion_count"] == 0

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 1.0  # all met: 2.0 / 2.0 = 1.0

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_retry_resolves_errored_criteria(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Batch retry resolves errored criteria with correct re-indexing."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            batch_recovery_enabled=False,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
            RubricItem(criterion="c3", weight=1.0),
        ]

        initial_verdicts = [
            Verdict(met=True, reasoning="ok"),
            Verdict(met=None, reasoning="timeout"),
            Verdict(met=None, reasoning="crash"),
        ]
        retry_verdicts = [
            Verdict(met=True, reasoning="ok retry"),
            Verdict(met=True, reasoning="ok retry 2"),
        ]
        mock_eval_all.side_effect = [
            (initial_verdicts, LLMUsage(cost_usd=0.1)),
            (retry_verdicts, LLMUsage(cost_usd=0.05)),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert all(r["met"] is True for r in info["criterion_results"])
        assert info["errored_criterion_count"] == 0

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 1.0  # all met: 3.0 / 3.0 = 1.0

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_retry_uses_split_scheduler(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            auto_parallel_target_chunk_size=1,
            rate_limit_discovery=False,
            batch_recovery_enabled=False,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c0", weight=1.0),
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
        ]
        retry_chunk_sizes: list[int] = []

        def _side_effect(
            judge_input: BatchJudgeInput,
            *,
            trace_path: str,
            **_kwargs: Any,
        ) -> tuple[list[Verdict], LLMUsage]:
            criterion = judge_input.criteria[0]
            if "_retry1" in trace_path:
                retry_chunk_sizes.append(len(judge_input.criteria))
                return [Verdict(met=True, reasoning=f"{criterion} retry ok")], LLMUsage()
            if criterion == "c0":
                return [Verdict(met=True, reasoning="ok")], LLMUsage()
            return [Verdict(met=None, reasoning="timeout")], LLMUsage()

        mock_eval_all.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert all(result["met"] is True for result in info["criterion_results"])
        assert retry_chunk_sizes == [1, 1]

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_recovery_preserves_successes_and_reruns_missing_verdicts(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            auto_parallel=False,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c0", weight=1.0),
            RubricItem(criterion="c1", weight=1.0),
        ]
        trace_paths: list[str] = []

        def _side_effect(
            judge_input: BatchJudgeInput,
            *,
            trace_path: str,
            **_kwargs: Any,
        ) -> tuple[list[Verdict], LLMUsage]:
            trace_name = os.path.basename(trace_path)
            trace_paths.append(trace_name)
            if "recovery" in trace_name:
                assert judge_input.criteria == ["c1"]
                return [Verdict(met=True, reasoning="recovered")], LLMUsage(cost_usd=0.02)
            return [Verdict(met=True, reasoning="ok")], LLMUsage(cost_usd=0.1)

        mock_eval_all.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert [result["met"] for result in info["criterion_results"]] == [True, True]
        assert info["batch_recovery"]["attempted_mini_batches"] == 1
        assert info["batch_recovery"]["attempted_criteria"] == 1
        assert info["batch_recovery"]["recovered_criteria"] == 1
        assert info["batch_recovery"]["extra_usage"]["cost_usd"] == pytest.approx(0.02)
        assert info["llm_usage"]["cost_usd"] == pytest.approx(0.12)
        assert trace_paths == ["judge_trace_batch.txt", "judge_trace_batch_recovery1_0.txt"]

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_recovery_splits_empty_batch_into_mini_batches(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            auto_parallel=False,
            batch_recovery_target_chunk_size=2,
        )
        mock_rubric.return_value = [RubricItem(criterion=f"c{i}", weight=1.0) for i in range(5)]

        def _side_effect(judge_input: BatchJudgeInput, *, trace_path: str, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if "recovery" not in os.path.basename(trace_path):
                return [
                    Verdict(met=None, reasoning="Judge agent wrote an empty verdict file.")
                    for _ in judge_input.criteria
                ], LLMUsage()
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()

        mock_eval_all.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert all(result["met"] is True for result in info["criterion_results"])
        assert info["batch_recovery"]["attempted_mini_batches"] == 3
        assert info["batch_recovery"]["attempted_criteria"] == 5
        assert info["batch_recovery"]["recovered_criteria"] == 5
        assert mock_eval_all.call_count == 4

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_recovery_ignores_rate_limit_failures(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
            auto_parallel=False,
        )
        mock_rubric.return_value = [RubricItem(criterion="c0", weight=1.0)]
        mock_eval_all.return_value = ([Verdict(met=None, reasoning="429 rate limit")], LLMUsage())

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()

        assert exc_info.value.code == 1
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert "batch_recovery" not in info
        assert mock_eval_all.call_count == 1

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_judge_retries_zero_disables_batch_recovery(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
            auto_parallel=False,
            batch_recovery_enabled=True,
        )
        mock_rubric.return_value = [RubricItem(criterion="c0", weight=1.0)]
        mock_eval_all.return_value = (
            [Verdict(met=None, reasoning="Judge agent wrote an empty verdict file.")],
            LLMUsage(),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()

        assert exc_info.value.code == 1
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert "batch_recovery" not in info
        assert info["errored_criterion_count"] == 1
        assert mock_eval_all.call_count == 1
        assert not (tmp_path / "output" / "reward.json").exists()

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_recovery_respects_max_rounds(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=2,
            mode="batch",
            auto_parallel=False,
            batch_recovery_max_rounds=2,
        )
        mock_rubric.return_value = [RubricItem(criterion="c0", weight=1.0)]
        mock_eval_all.return_value = (
            [Verdict(met=None, reasoning="Judge agent wrote an empty verdict file.")],
            LLMUsage(),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()

        assert exc_info.value.code == 1
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["batch_recovery"]["rounds"] == 2
        assert info["batch_recovery"]["attempted_mini_batches"] == 2
        assert info["batch_recovery"]["unresolved_criteria"] == 1
        assert info["errored_criterion_count"] == 1

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_judge_retries_zero_disables_retry(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """judge_retries=0 skips retry loop entirely — errors cause hard fail."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="individual",
        )
        mock_rubric.return_value = [RubricItem(criterion="c1", weight=1.0)]
        mock_eval.return_value = ([Verdict(met=None, reasoning="timeout")], LLMUsage())

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()
            assert exc_info.value.code == 1

        assert (tmp_path / "output" / "info.json").exists()
        assert not (tmp_path / "output" / "reward.json").exists()
        assert mock_eval.call_count == 1

    @patch("gandalf.orchestrator.time.sleep")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_persistent_rate_limit_writes_info_not_reward(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_sleep: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
            auto_parallel_target_chunk_size=1,
            max_concurrency=2,
            rate_limit_backoff_seconds=1,
            rate_limit_max_requeues_per_chunk=1,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            return [Verdict(met=None, reasoning="429 rate limit") for _ in judge_input.criteria], LLMUsage()

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()
            assert exc_info.value.code == 1

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["errored_criterion_count"] == 2
        assert all(result["met"] is None for result in info["criterion_results"])
        assert mock_eval.call_count == 4
        assert not (tmp_path / "output" / "reward.json").exists()

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_hard_fail_writes_info_not_reward(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Persistent errors: info.json written, reward.json NOT written, exit 1."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="individual",
        )
        mock_rubric.return_value = [RubricItem(criterion="c1", weight=1.0)]
        mock_eval.return_value = ([Verdict(met=None, reasoning="always fails")], LLMUsage())

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()
            assert exc_info.value.code == 1

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["criterion_results"][0]["met"] is None
        assert info["errored_criterion_count"] == 1
        assert not (tmp_path / "output" / "reward.json").exists()

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_all_resolved_after_retry(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """After retry resolves all errors: reward.json written with correct reward."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="individual",
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
        ]

        mock_eval.side_effect = [
            ([Verdict(met=True, reasoning="ok")], LLMUsage()),
            ([Verdict(met=None, reasoning="timeout")], LLMUsage()),
            ([Verdict(met=False, reasoning="genuinely failed")], LLMUsage()),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 0.5  # c1 met, c2 not: 1.0 / 2.0 = 0.5

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["errored_criterion_count"] == 0

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_reward_json_with_negative_weights(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """reward.json must contain the [0,1] reward, not the raw score,
        when negative-weight criteria are present."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="individual",
        )
        mock_rubric.return_value = [
            RubricItem(criterion="correct output", weight=3.0),
            RubricItem(criterion="used hardcoded values", weight=-1.0),
        ]

        # Both criteria met: raw = 3 + (-1) = 2, reward = 2/3 ≈ 0.6667
        mock_eval.side_effect = [
            ([Verdict(met=True, reasoning="ok")], LLMUsage()),
            ([Verdict(met=True, reasoning="hardcoded detected")], LLMUsage()),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        info = json.loads((tmp_path / "output" / "info.json").read_text())

        assert reward["reward"] == 0.6667
        assert info["raw_score"] == 2.0
        assert info["reward"] == 0.6667
        assert reward["reward"] == info["reward"]

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_failed_section_gate_skips_section_criteria(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """A failed section-level gate skips that section's criteria before judging."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                section_gate="The submitted workbook exists.",
            ),
            RubricItem(criterion="Used hardcoded values", weight=-2.0, section="Model Build"),
            RubricItem(criterion="Includes README", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if mock_eval.call_count == 1:
                assert judge_input.criteria == ["The submitted workbook exists."]
                return [Verdict(met=False, reasoning="Workbook missing")], LLMUsage(cost_usd=0.01)
            assert judge_input.criteria == ["Includes README"]
            return [Verdict(met=True, reasoning="ok")], LLMUsage(cost_usd=0.02)

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        assert mock_eval.call_count == 2
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["raw_score"] == 1.0
        assert info["evaluated_criteria_pct"] == 33.33
        assert info["criterion_results"][0]["skipped"] is True
        assert info["criterion_results"][1]["skipped"] is True
        assert info["criterion_results"][1]["score_contribution"] == 0.0
        assert info["section_results"][0]["section_gate_met"] is False
        assert info["section_results"][0]["score"] == 0.0
        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 0.1667

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_multiple_section_gates_pass_before_criteria(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                section_gates=["The submitted workbook exists.", "The submitted workbook opens."],
            ),
            RubricItem(criterion="Includes README", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if mock_eval.call_count == 1:
                assert judge_input.criteria == ["The submitted workbook exists."]
                return [Verdict(met=True, reasoning="exists")], LLMUsage(cost_usd=0.01)
            if mock_eval.call_count == 2:
                assert judge_input.criteria == ["The submitted workbook opens."]
                return [Verdict(met=True, reasoning="opens")], LLMUsage(cost_usd=0.01)
            assert judge_input.criteria == ["Forecast formulas are correct", "Includes README"]
            return [
                Verdict(met=True, reasoning="ok"),
                Verdict(met=True, reasoning="ok"),
            ], LLMUsage(cost_usd=0.02)

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        assert mock_eval.call_count == 3
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        section_result = info["section_results"][0]
        assert info["raw_score"] == 6.0
        assert section_result["section_gate_count"] == 2
        assert section_result["passed_section_gate_indices"] == [0, 1]
        assert section_result["section_gates_met"] is True
        assert section_result["gated_out"] is False

    @patch("gandalf.orchestrator.run_judge")
    def test_section_gate_round_splits_independent_sections(
        self,
        mock_eval: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)
        config = make_config(
            workdir=str(tmp_path),
            output_dir=output_dir,
            mode="batch",
            auto_parallel_target_chunk_size=2,
            max_concurrency=2,
            rate_limit_discovery=False,
        )
        section_gates = {f"Section {idx}": [f"Gate {idx}"] for idx in range(5)}
        chunk_sizes: list[int] = []
        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            chunk_sizes.append(len(judge_input.criteria))
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)
            with lock:
                current_concurrent -= 1
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()

        mock_eval.side_effect = _side_effect

        results, _usage, initial_errored = evaluate_section_gates(
            config,
            section_gates,
            "done",
            "test",
            "",
            None,
        )

        assert sorted(chunk_sizes) == [1, 2, 2]
        assert peak_concurrent == 2
        assert initial_errored == 0
        assert all(gate_result.met is True for gate_results in results.values() for gate_result in gate_results)

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_first_section_gate_failure_skips_later_gates_and_criteria(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                section_gates=["The submitted workbook exists.", "The submitted workbook opens."],
            ),
            RubricItem(criterion="Includes README", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if mock_eval.call_count == 1:
                assert judge_input.criteria == ["The submitted workbook exists."]
                return [Verdict(met=False, reasoning="missing")], LLMUsage(cost_usd=0.01)
            assert judge_input.criteria == ["Includes README"]
            return [Verdict(met=True, reasoning="ok")], LLMUsage(cost_usd=0.02)

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        assert mock_eval.call_count == 2
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        section_result = info["section_results"][0]
        assert info["errored_section_gate_count"] == 0
        assert section_result["failed_section_gate_indices"] == [0]
        assert section_result["skipped_section_gate_indices"] == [1]
        assert section_result["section_gate_results"][1]["skipped"] is True
        assert section_result["section_gates_met"] is False
        assert info["criterion_results"][0]["skipped"] is True

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_later_section_gate_failure_skips_criteria(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=0,
            mode="batch",
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                section_gates=["The submitted workbook exists.", "The submitted workbook opens."],
            ),
            RubricItem(criterion="Includes README", weight=1.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if mock_eval.call_count == 1:
                assert judge_input.criteria == ["The submitted workbook exists."]
                return [Verdict(met=True, reasoning="exists")], LLMUsage(cost_usd=0.01)
            if mock_eval.call_count == 2:
                assert judge_input.criteria == ["The submitted workbook opens."]
                return [Verdict(met=False, reasoning="does not open")], LLMUsage(cost_usd=0.01)
            assert judge_input.criteria == ["Includes README"]
            return [Verdict(met=True, reasoning="ok")], LLMUsage(cost_usd=0.02)

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        assert mock_eval.call_count == 3
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        section_result = info["section_results"][0]
        assert section_result["passed_section_gate_indices"] == [0]
        assert section_result["failed_section_gate_indices"] == [1]
        assert section_result["skipped_section_gate_indices"] == []
        assert info["criterion_results"][0]["skipped"] is True

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_errored_section_gate_hard_fails_after_retries(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            batch_recovery_enabled=False,
        )
        mock_rubric.return_value = [
            RubricItem(
                criterion="Forecast formulas are correct",
                weight=5.0,
                section="Model Build",
                section_gates=["The submitted workbook exists.", "The submitted workbook opens."],
            )
        ]
        mock_eval.return_value = ([Verdict(met=None, reasoning="judge crashed")], LLMUsage())

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            with pytest.raises(SystemExit) as exc_info:
                main()
            assert exc_info.value.code == 1

        assert mock_eval.call_count == 2
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["errored_section_gate_count"] == 1
        assert info["section_results"][0]["errored_section_gate_indices"] == [0]
        assert info["section_results"][0]["skipped_section_gate_indices"] == [1]
        assert info["section_results"][0]["section_gate_results"][1]["skipped"] is True
        assert info["criterion_results"][0]["skipped"] is True
        assert not (tmp_path / "output" / "reward.json").exists()


class TestCloneWorkspace:
    """Tests for clone_workspace resilience to unreadable files."""

    def test_readable_files_are_cloned(self, tmp_path: pathlib.Path) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "file.txt").write_text("hello")
        (workspace / "subdir").mkdir()
        (workspace / "subdir" / "nested.txt").write_text("world")

        clone_dir = clone_workspace(str(workspace))
        try:
            assert (pathlib.Path(clone_dir) / "file.txt").read_text() == "hello"
            assert (pathlib.Path(clone_dir) / "subdir" / "nested.txt").read_text() == "world"
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_unreadable_files_are_skipped_not_fatal(self, tmp_path: pathlib.Path) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "readable.txt").write_text("ok")

        restricted = workspace / "restricted.txt"
        restricted.write_text("secret")
        restricted.chmod(0o000)

        try:
            clone_dir = clone_workspace(str(workspace))
            cloned = pathlib.Path(clone_dir)
            assert (cloned / "readable.txt").read_text() == "ok"
            assert not (cloned / "restricted.txt").exists()
        finally:
            restricted.chmod(0o644)
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_skipped_files_are_logged(self, tmp_path: pathlib.Path, capsys: pytest.CaptureFixture[str]) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        restricted = workspace / "noperm.txt"
        restricted.write_text("x")
        restricted.chmod(0o000)

        try:
            clone_dir = clone_workspace(str(workspace))
            stderr = capsys.readouterr().err
            assert "skipped 1 unreadable path(s)" in stderr
            assert "noperm.txt" in stderr
        finally:
            restricted.chmod(0o644)
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_unreadable_directory_is_skipped_not_fatal(self, tmp_path: pathlib.Path) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "readable.txt").write_text("ok")

        # Create a directory tree and make the parent unreadable
        restricted_dir = workspace / ".tool_cache"
        restricted_dir.mkdir()
        (restricted_dir / "data.bin").write_text("cached")
        restricted_dir.chmod(0o000)

        try:
            clone_dir = clone_workspace(str(workspace))
            cloned = pathlib.Path(clone_dir)
            assert (cloned / "readable.txt").read_text() == "ok"
            # The restricted directory's contents should not appear
            assert not (cloned / ".tool_cache" / "data.bin").exists()
        finally:
            restricted_dir.chmod(0o755)
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_unreadable_directory_is_logged(self, tmp_path: pathlib.Path, capsys: pytest.CaptureFixture[str]) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        restricted_dir = workspace / ".cache"
        restricted_dir.mkdir()
        restricted_dir.chmod(0o000)

        try:
            clone_dir = clone_workspace(str(workspace))
            stderr = capsys.readouterr().err
            assert "skipped 1 unreadable path(s)" in stderr
            assert ".cache" in stderr
        finally:
            restricted_dir.chmod(0o755)
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_clone_is_group_writable(self, tmp_path: pathlib.Path) -> None:
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "file.txt").write_text("data")

        clone_dir = clone_workspace(str(workspace))
        try:
            cloned = pathlib.Path(clone_dir)
            assert os.stat(clone_dir).st_mode & 0o070 == 0o070
            fstat = os.stat(cloned / "file.txt")
            assert fstat.st_mode & 0o060 == 0o060
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_clone_workspace_uses_parent_dir_when_provided(self, tmp_path: pathlib.Path) -> None:
        workspace = tmp_path / "workspace"
        parent = tmp_path / "output" / "judge_workspaces"
        workspace.mkdir()
        (workspace / "file.txt").write_text("data")

        clone_dir = clone_workspace(str(workspace), parent_dir=str(parent))
        try:
            clone = pathlib.Path(clone_dir)
            assert clone.parent == parent
            assert (clone / "file.txt").read_text() == "data"
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_judge_clone_parent_dir_uses_excel_access_for_mac_excel(
        self, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr("gandalf.orchestrator.excel_access_root", lambda: tmp_path / "excel_access")
        config = make_config(
            output_dir=str(tmp_path / "output"),
            excel_backend={"enabled": True, "backend": "mac_excel"},
        )

        # Judge clones go under the stable excel_access root, not the per-run output dir.
        assert judge_clone_parent_dir(config) == str(tmp_path / "excel_access" / "judge_workspaces")

    def test_judge_clone_parent_dir_uses_default_temp_for_non_mac_excel(self, tmp_path: pathlib.Path) -> None:
        config = make_config(
            output_dir=str(tmp_path / "output"),
            excel_backend={"enabled": False},
        )

        assert judge_clone_parent_dir(config) is None

    def test_clone_is_world_accessible(self, tmp_path: pathlib.Path) -> None:
        """Clone dir must have world execute+write so sandbox_user can use it.

        Regression: shutil.copytree preserved the source workspace permissions
        (typically world-executable) on the root clone dir.  The new os.walk
        implementation creates clone_dir via mkdtemp (mode 0o700) and must
        explicitly grant world bits — otherwise sandbox_user (not in the
        grader's group) cannot traverse or write to the workspace.

        This test fails on the pre-fix code (|0o070 → 0o770, no world bits)
        and passes with the fix (|0o077 → 0o777).
        """
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "file.txt").write_text("hello")
        (workspace / "subdir").mkdir()
        (workspace / "subdir" / "nested.txt").write_text("world")

        clone_dir = clone_workspace(str(workspace))
        try:
            clone = pathlib.Path(clone_dir)

            # Root clone dir: world execute (traverse) + write (create files inside it)
            root_mode = clone.stat().st_mode
            assert root_mode & 0o001, (
                "clone root missing world execute — sandbox_user cannot traverse it "
                "(regression: os.walk+mkdtemp loses the world-execute bit that "
                "shutil.copytree preserved from the source workspace)"
            )
            assert root_mode & 0o002, "clone root missing world write — sandbox_user cannot create files in it"

            # Subdirectories must also have world execute+write
            sub_mode = (clone / "subdir").stat().st_mode
            assert sub_mode & 0o001, "subdir missing world execute"
            assert sub_mode & 0o002, "subdir missing world write"

            # Files must have world read so sandbox_user can inspect them
            file_mode = (clone / "file.txt").stat().st_mode
            assert file_mode & 0o004, "file missing world read"
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_executable_bits_are_preserved(self, tmp_path: pathlib.Path) -> None:
        """Cloned files must retain execute bits so scripts/binaries remain runnable."""
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        script = workspace / "run.sh"
        script.write_text("#!/bin/sh\necho hi")
        script.chmod(0o755)
        data = workspace / "data.txt"
        data.write_text("plain")

        clone_dir = clone_workspace(str(workspace))
        try:
            cloned = pathlib.Path(clone_dir)
            cloned_script_mode = (cloned / "run.sh").stat().st_mode
            assert cloned_script_mode & 0o111, (
                f"Executable bits lost on cloned script (mode={oct(cloned_script_mode)}) — "
                "judge runs that execute workspace scripts will break"
            )
            # Non-executable file should NOT gain execute bits
            cloned_data_mode = (cloned / "data.txt").stat().st_mode
            assert not (cloned_data_mode & 0o111), (
                f"Non-executable file gained execute bits (mode={oct(cloned_data_mode)})"
            )
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_broken_symlink_is_skipped_not_fatal(self, tmp_path: pathlib.Path) -> None:
        """A broken symlink in the workspace must be skipped, not crash the clone.

        The old code caught only PermissionError; shutil.copy2 on a broken
        symlink raises FileNotFoundError (an OSError subclass), which would
        have propagated and aborted the entire clone.
        """
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "good.txt").write_text("ok")
        (workspace / "broken_link").symlink_to("/nonexistent/target")

        clone_dir = clone_workspace(str(workspace))
        try:
            cloned = pathlib.Path(clone_dir)
            assert (cloned / "good.txt").read_text() == "ok"
            assert not (cloned / "broken_link").exists()
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_symlink_to_directory_is_skipped_not_fatal(self, tmp_path: pathlib.Path) -> None:
        """A symlink-to-directory in filenames must be skipped, not crash the clone.

        os.walk (followlinks=False) places dir-symlinks in filenames.
        shutil.copy2 on them raises IsADirectoryError (OSError subclass).
        """
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        real_dir = tmp_path / "real_dir"
        real_dir.mkdir()
        (real_dir / "data.txt").write_text("data")
        (workspace / "good.txt").write_text("ok")
        (workspace / "dir_link").symlink_to(real_dir)

        clone_dir = clone_workspace(str(workspace))
        try:
            cloned = pathlib.Path(clone_dir)
            assert (cloned / "good.txt").read_text() == "ok"
            # The symlink itself should not have been copied as a file
            assert not (cloned / "dir_link").is_file()
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)

    def test_symlink_loop_is_skipped_not_fatal(self, tmp_path: pathlib.Path) -> None:
        """Circular symlinks in the workspace must be skipped, not crash the clone."""
        workspace = tmp_path / "workspace"
        workspace.mkdir()
        (workspace / "good.txt").write_text("ok")
        # Create a symlink loop: a -> b -> a
        (workspace / "loop_a").symlink_to(workspace / "loop_b")
        (workspace / "loop_b").symlink_to(workspace / "loop_a")

        clone_dir = clone_workspace(str(workspace))
        try:
            cloned = pathlib.Path(clone_dir)
            assert (cloned / "good.txt").read_text() == "ok"
            assert not (cloned / "loop_a").exists()
            assert not (cloned / "loop_b").exists()
        finally:
            shutil.rmtree(clone_dir, ignore_errors=True)


class TestBatchRetryNegativeWeights:
    """Issue 9: batch mode + retries + negative weights combined."""

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_retry_with_negative_weights(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Batch retry with negative weights: errored negative-weight criterion
        resolves on retry and the penalty is correctly applied."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
            batch_recovery_enabled=False,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="correct output", weight=3.0),
            RubricItem(criterion="no hardcoded values", weight=-1.0),
            RubricItem(criterion="has tests", weight=2.0),
        ]

        # Initial: c0 passes, c1 (negative) errors, c2 passes
        initial_verdicts = [
            Verdict(met=True, reasoning="ok"),
            Verdict(met=None, reasoning="timeout"),
            Verdict(met=True, reasoning="ok"),
        ]
        # Retry: c1 (negative) met — penalty applies
        retry_verdicts = [
            Verdict(met=True, reasoning="hardcoded detected"),
        ]
        mock_eval_all.side_effect = [
            (initial_verdicts, LLMUsage(cost_usd=0.1)),
            (retry_verdicts, LLMUsage(cost_usd=0.02)),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["errored_criterion_count"] == 0
        # raw = 3 + (-1) + 2 = 4, max_positive = 5, reward = 4/5 = 0.8
        assert info["raw_score"] == 4.0
        assert info["reward"] == 0.8

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 0.8


class TestRetryJudgePromptPassthrough:
    """Issue 10: verify resolve_judge_prompt flows through to retry calls."""

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value="CUSTOM {{ criterion }}")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_sequential_retry_passes_judge_prompt(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Custom judge_prompt must be forwarded to retry evaluate_criterion calls."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="individual",
        )
        mock_rubric.return_value = [RubricItem(criterion="c1", weight=1.0)]

        # First call errors, retry succeeds
        mock_eval.side_effect = [
            ([Verdict(met=None, reasoning="timeout")], LLMUsage()),
            ([Verdict(met=True, reasoning="ok")], LLMUsage()),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        # Both the initial call and the retry call should have judge_prompt set
        assert mock_eval.call_count == 2
        for call in mock_eval.call_args_list:
            judge_input = call[0][0]
            assert judge_input.judge_prompt == "CUSTOM {{ criterion }}"

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value="BATCH CUSTOM {{ criteria }}")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_batch_retry_passes_judge_prompt(
        self,
        mock_eval_all: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """Custom judge_prompt must be forwarded to retry evaluate_all_criteria calls."""

        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            judge_retries=1,
            mode="batch",
        )
        mock_rubric.return_value = [
            RubricItem(criterion="c1", weight=1.0),
            RubricItem(criterion="c2", weight=1.0),
        ]

        mock_eval_all.side_effect = [
            ([Verdict(met=True, reasoning="ok"), Verdict(met=None, reasoning="err")], LLMUsage()),
            ([Verdict(met=True, reasoning="ok retry")], LLMUsage()),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        assert mock_eval_all.call_count == 2
        for call in mock_eval_all.call_args_list:
            judge_input = call[0][0]
            assert judge_input.judge_prompt == "BATCH CUSTOM {{ criteria }}"


class TestGatesOnlyCli:
    """Tests for the CLI-only gate preflight mode."""

    def _config(self, tmp_path: pathlib.Path, *, judge_retries: int = 0) -> GraderConfig:
        return make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            auto_parallel=False,
            judge_retries=judge_retries,
            batch_recovery_enabled=False,
        )

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    @patch("gandalf.orchestrator.run_workbook_repair_check")
    def test_gates_only_evaluates_section_gates_and_gate_criteria_only(
        self,
        mock_repair_check: Any,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path)
        output_dir = tmp_path / "output"
        output_dir.mkdir()
        (output_dir / "reward.json").write_text('{"reward": 1.0}')
        mock_rubric.return_value = [
            RubricItem(
                criterion="Criterion gate",
                weight=0.0,
                section="Section A",
                gate=True,
                section_gates=["Section gate"],
            ),
            RubricItem(criterion="Section non-gate", weight=5.0, section="Section A"),
            RubricItem(criterion="Top-level non-gate", weight=2.0),
        ]

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            if mock_eval.call_count == 1:
                assert judge_input.criteria == ["Section gate"]
                return [Verdict(met=True, reasoning="section gate passed")], LLMUsage(cost_usd=0.01)
            assert judge_input.criteria == ["Criterion gate"]
            return [Verdict(met=True, reasoning="criterion gate passed")], LLMUsage(cost_usd=0.02)

        mock_eval.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml", "--gates-only"]):
            main()

        mock_repair_check.assert_not_called()
        assert mock_eval.call_count == 2
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert "workbook_repair_check" not in info
        assert info["evaluated_criteria_pct"] == 33.33
        assert info["criterion_results"][0]["gate"] is True
        assert info["criterion_results"][0]["met"] is True
        assert info["criterion_results"][1]["skipped"] is True
        assert info["criterion_results"][1]["skip_reason"] == "Skipped because --gates-only was set."
        assert info["criterion_results"][2]["skipped"] is True
        assert info["section_results"][0]["passed_section_gate_indices"] == [0]
        assert not (tmp_path / "output" / "reward.json").exists()
        assert "Gate check passed: 2/2 gate checks passed." in capsys.readouterr().out

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_gates_only_exits_one_when_gate_criterion_fails(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path)
        mock_rubric.return_value = [
            RubricItem(criterion="Criterion gate", weight=0.0, section="Section A", gate=True),
            RubricItem(criterion="Non-gate", weight=5.0, section="Section A"),
        ]
        mock_run_batch.return_value = (
            [
                CriterionResult(
                    criterion="Criterion gate",
                    weight=0.0,
                    section="Section A",
                    gate=True,
                    met=False,
                    reasoning="failed",
                )
            ],
            LLMUsage(),
        )

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--gates-only"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 1
        called_rubric = mock_run_batch.call_args[0][1]
        assert [item.criterion for item in called_rubric] == ["Criterion gate"]
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["criterion_results"][0]["met"] is False
        assert info["criterion_results"][1]["skipped"] is True
        assert not (tmp_path / "output" / "reward.json").exists()
        assert "Gate check failed: 1 failed, 0 errored gate check(s)." in capsys.readouterr().err

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_gates_only_section_gate_failure_skips_later_gates_and_criteria(
        self,
        mock_eval: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path)
        mock_rubric.return_value = [
            RubricItem(
                criterion="Criterion gate",
                weight=0.0,
                section="Section A",
                gate=True,
                section_gates=["First section gate", "Second section gate"],
            ),
            RubricItem(criterion="Section non-gate", weight=5.0, section="Section A"),
        ]
        mock_eval.return_value = ([Verdict(met=False, reasoning="section gate failed")], LLMUsage())

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--gates-only"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 1
        assert mock_eval.call_count == 1
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        section_result = info["section_results"][0]
        assert section_result["failed_section_gate_indices"] == [0]
        assert section_result["skipped_section_gate_indices"] == [1]
        assert info["criterion_results"][0]["skipped"] is True
        assert info["criterion_results"][1]["skipped"] is True
        assert not (tmp_path / "output" / "reward.json").exists()
        assert "Gate check failed: 1 failed, 0 errored gate check(s)." in capsys.readouterr().err

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_gates_only_retries_errored_gate_criteria_and_exits_one(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path, judge_retries=1)
        mock_rubric.return_value = [
            RubricItem(criterion="Criterion gate", weight=0.0, section="Section A", gate=True),
            RubricItem(criterion="Non-gate", weight=5.0, section="Section A"),
        ]
        mock_run_batch.side_effect = [
            (
                [
                    CriterionResult(
                        criterion="Criterion gate",
                        weight=0.0,
                        section="Section A",
                        gate=True,
                        met=None,
                        reasoning="timeout",
                    )
                ],
                LLMUsage(),
            ),
            (
                [
                    CriterionResult(
                        criterion="Criterion gate",
                        weight=0.0,
                        section="Section A",
                        gate=True,
                        met=None,
                        reasoning="timeout",
                    )
                ],
                LLMUsage(),
            ),
        ]

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--gates-only"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 1
        assert mock_run_batch.call_count == 2
        assert mock_run_batch.call_args_list[1].kwargs["trace_suffix"] == "_retry1"
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["errored_criterion_count"] == 1
        assert info["criterion_results"][1]["skipped"] is True
        assert not (tmp_path / "output" / "reward.json").exists()
        assert "Gate check failed: 0 failed, 1 errored gate check(s)." in capsys.readouterr().err

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_gates_only_no_gates_exits_zero_without_reward(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path)
        mock_rubric.return_value = [
            RubricItem(criterion="Non-gate A", weight=2.0),
            RubricItem(criterion="Non-gate B", weight=3.0),
        ]

        with patch("sys.argv", ["prog", "--config", "dummy.toml", "--gates-only"]):
            main()

        mock_run_batch.assert_not_called()
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["evaluated_criteria_pct"] == 0.0
        assert all(result["skipped"] for result in info["criterion_results"])
        assert not (tmp_path / "output" / "reward.json").exists()
        assert "No gates to evaluate." in capsys.readouterr().out

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_gates_only_respects_section_filters_before_gate_selection(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = self._config(tmp_path)
        mock_rubric.return_value = [
            RubricItem(criterion="A gate", weight=0.0, section="Section A", gate=True),
            RubricItem(criterion="B gate", weight=0.0, section="Section B", gate=True),
            RubricItem(criterion="Top gate", weight=0.0, gate=True),
        ]
        mock_run_batch.return_value = (
            [CriterionResult(criterion="B gate", weight=0.0, section="Section B", gate=True, met=True, reasoning="ok")],
            LLMUsage(),
        )

        with patch(
            "sys.argv",
            ["prog", "--config", "dummy.toml", "--gates-only", "--include-section", "Section B"],
        ):
            main()

        called_rubric = mock_run_batch.call_args[0][1]
        assert [item.criterion for item in called_rubric] == ["B gate"]
        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert [result["criterion"] for result in info["criterion_results"]] == ["B gate"]
        assert "Gate check passed: 1/1 gate checks passed." in capsys.readouterr().out


class TestSectionFilters:
    """Tests for CLI-only rubric section filtering."""

    def _rubric(self) -> list[RubricItem]:
        return [
            RubricItem(criterion="top", weight=5.0),
            RubricItem(criterion="a1", weight=2.0, section="Section A", section_gates=["A gate"]),
            RubricItem(criterion="a2", weight=1.0, section="Section A"),
            RubricItem(criterion="b1", weight=3.0, section="Section B", section_gates=["B gate"]),
            RubricItem(criterion="c1", weight=4.0, section="Section C"),
        ]

    def test_include_one_section(self) -> None:
        filtered = filter_rubric_sections(self._rubric(), include_sections=["Section A"])

        assert [item.criterion for item in filtered] == ["a1", "a2"]

    def test_include_multiple_sections_preserves_rubric_order(self) -> None:
        filtered = filter_rubric_sections(
            self._rubric(),
            include_sections=["Section C", "Section A"],
        )

        assert [item.criterion for item in filtered] == ["a1", "a2", "c1"]

    def test_exclude_section_keeps_unsectioned_criteria(self) -> None:
        filtered = filter_rubric_sections(self._rubric(), exclude_sections=["Section B"])

        assert [item.criterion for item in filtered] == ["top", "a1", "a2", "c1"]

    def test_include_then_exclude(self) -> None:
        filtered = filter_rubric_sections(
            self._rubric(),
            include_sections=["Section A", "Section B"],
            exclude_sections=["Section A"],
        )

        assert [item.criterion for item in filtered] == ["b1"]

    def test_unknown_section_lists_available_sections(self) -> None:
        with pytest.raises(ValueError, match="Unknown section filter value"):
            filter_rubric_sections(self._rubric(), include_sections=["Section Z"])

        try:
            filter_rubric_sections(self._rubric(), include_sections=["Section Z"])
        except ValueError as e:
            assert "Section Z" in str(e)
            assert "Section A, Section B, Section C" in str(e)

    def test_empty_filtered_rubric_raises(self) -> None:
        with pytest.raises(ValueError, match="left no rubric criteria"):
            filter_rubric_sections(
                self._rubric(),
                include_sections=["Section A"],
                exclude_sections=["Section A"],
            )

    def test_section_gate_collection_uses_only_retained_sections(self) -> None:
        filtered = filter_rubric_sections(self._rubric(), include_sections=["Section B"])

        assert collect_section_gates(filtered) == {"Section B": ["B gate"]}

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_include_section_filters_run_and_denominator(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        output_dir = str(tmp_path / "output")
        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="batch",
            auto_parallel=False,
        )
        mock_rubric.return_value = [
            RubricItem(criterion="a", weight=2.0, section="Section A"),
            RubricItem(criterion="b", weight=8.0, section="Section B"),
        ]

        def _side_effect(
            _config: GraderConfig,
            rubric: list[RubricItem],
            *_args: Any,
            **_kwargs: Any,
        ) -> tuple[list[CriterionResult], LLMUsage]:
            assert [item.criterion for item in rubric] == ["a"]
            return [
                CriterionResult(
                    criterion=rubric[0].criterion,
                    weight=rubric[0].weight,
                    section=rubric[0].section,
                    met=True,
                    reasoning="ok",
                )
            ], LLMUsage()

        mock_run_batch.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml", "--include-section", "Section A"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert info["maximum_score"] == 2.0
        assert info["raw_score"] == 2.0
        assert info["reward"] == 1.0
        assert [result["section"] for result in info["criterion_results"]] == ["Section A"]

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_unknown_section_errors_before_judge(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = make_config(workdir=str(tmp_path), output_dir=str(tmp_path / "output"))
        mock_rubric.return_value = [RubricItem(criterion="a", weight=1.0, section="Section A")]

        with (
            patch("sys.argv", ["prog", "--config", "dummy.toml", "--include-section", "Section Z"]),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 2
        assert "Unknown section filter value" in capsys.readouterr().err
        mock_trajectory.assert_not_called()
        mock_run_batch.assert_not_called()

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_empty_filter_errors_before_judge(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
        capsys: pytest.CaptureFixture[str],
    ) -> None:
        mock_config.return_value = make_config(workdir=str(tmp_path), output_dir=str(tmp_path / "output"))
        mock_rubric.return_value = [RubricItem(criterion="a", weight=1.0, section="Section A")]

        with (
            patch(
                "sys.argv",
                [
                    "prog",
                    "--config",
                    "dummy.toml",
                    "--include-section",
                    "Section A",
                    "--exclude-section",
                    "Section A",
                ],
            ),
            pytest.raises(SystemExit) as exc_info,
        ):
            main()

        assert exc_info.value.code == 2
        assert "left no rubric criteria" in capsys.readouterr().err
        mock_trajectory.assert_not_called()
        mock_run_batch.assert_not_called()


class TestBatchConcurrent:
    """Tests for run_batch_concurrent — parallel positional splitting of batch evaluation."""

    def _make_rubric(self, n: int) -> list[RubricItem]:
        return [RubricItem(criterion=f"criterion {i}", weight=1.0) for i in range(n)]

    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_auto_parallel_false_dispatches_to_run_batch(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        """auto_parallel=false preserves the old single-batch dispatch."""
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="batch",
            auto_parallel=False,
        )
        rubric = self._make_rubric(2)
        mock_rubric.return_value = rubric

        mock_run_batch.return_value = (
            [
                CriterionResult(criterion="criterion 0", weight=1.0, met=True, reasoning="ok"),
                CriterionResult(criterion="criterion 1", weight=1.0, met=True, reasoning="ok"),
            ],
            LLMUsage(cost_usd=0.1),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        mock_run_batch.assert_called_once()

    def test_empty_rubric(self, tmp_path: pathlib.Path) -> None:
        """Empty rubric returns empty results without crashing."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)

        results, usage = run_batch_concurrent(config, [], "done", "test", "", None)

        assert results == []
        assert usage == LLMUsage()

    @patch("gandalf.orchestrator.run_judge")
    def test_splits_2_even(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """4 criteria split into 2 chunks of 2, results merged in order."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            verdicts = [Verdict(met=True, reasoning=f"ok {i}") for i in range(len(judge_input.criteria))]
            usage = LLMUsage(cost_usd=0.1, prompt_tokens=100, completion_tokens=50, cache_read_tokens=10)
            return verdicts, usage

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert len(results) == 4
        # Verify order preserved
        for i, r in enumerate(results):
            assert r.criterion == f"criterion {i}"
            assert r.met is True

        # 2 splits, each with usage
        assert usage.cost_usd == pytest.approx(0.2)
        assert usage.prompt_tokens == 200
        assert usage.completion_tokens == 100
        assert usage.cache_read_tokens == 20

        # Verify run_judge was called twice (one per split)
        assert mock_run_judge.call_count == 2

    @patch("gandalf.orchestrator.run_judge")
    def test_split_uses_local_indices(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Chunks must use 0-based local indices, not global rubric positions.

        Regression test: the judge prompt says "0 through N-1" and
        read_batch_verdict filters by 0 <= idx < N, so passing global
        indices (e.g. 3, 4, 5) for chunk 2 causes the judge to either
        write mismatched indices or have its verdicts silently discarded.
        """
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=3,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(6)  # 6 criteria → 3 chunks of 2

        received_criteria_counts: list[int] = []

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            received_criteria_counts.append(len(judge_input.criteria))
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage(cost_usd=0.05)

        mock_run_judge.side_effect = _side_effect

        results, _ = run_batch_concurrent(config, rubric, "done", "test", "", None)

        # Each chunk gets its own local criteria list
        assert all(c == 2 for c in received_criteria_counts)

        # All 6 results should be successful
        assert len(results) == 6
        assert all(r.met is True for r in results)
        # Results are in original rubric order
        for i, r in enumerate(results):
            assert r.criterion == f"criterion {i}"

    @patch("gandalf.orchestrator.run_judge")
    def test_splits_3_uneven(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """7 criteria split into chunks of [3, 3, 1]."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=3,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(7)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            verdicts = [Verdict(met=True, reasoning=f"ok {i}") for i in range(len(judge_input.criteria))]
            return verdicts, LLMUsage(cost_usd=0.1)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert len(results) == 7
        for i, r in enumerate(results):
            assert r.criterion == f"criterion {i}"

        assert mock_run_judge.call_count == 3
        assert usage.cost_usd == pytest.approx(0.3)

    @patch("gandalf.orchestrator.run_judge")
    def test_auto_splits_66_criteria_into_3_chunks(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Default batch auto-parallelism splits by the configured target chunk size."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(66)
        chunk_sizes: list[int] = []

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            chunk_sizes.append(len(judge_input.criteria))
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage(cost_usd=0.1)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert chunk_sizes == [22, 22, 22]
        assert len(results) == 66
        assert all(result.met is True for result in results)
        assert usage.cost_usd == pytest.approx(0.3)

    @patch("gandalf.orchestrator.run_judge")
    def test_adaptive_executor_ramps_concurrency(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Adaptive discovery starts at 2 and ramps up to the configured cap."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=8,
            max_concurrency=4,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(8)
        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)
            with lock:
                current_concurrent -= 1
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()

        mock_run_judge.side_effect = _side_effect

        results, _usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert len(results) == 8
        assert peak_concurrent == 4

    @patch("gandalf.orchestrator.time.sleep")
    @patch("gandalf.orchestrator.run_judge")
    def test_rate_limited_chunks_are_requeued_and_concurrency_backs_off(
        self,
        mock_run_judge: Any,
        mock_sleep: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        """A 429 chunk is requeued and later chunks run at the last successful concurrency."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=8,
            max_concurrency=4,
            rate_limit_backoff_seconds=1,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(8)
        attempts: dict[int, int] = {}
        later_peak_concurrent = 0
        later_current_concurrent = 0
        lock = threading.Lock()

        def _first_idx(judge_input: BatchJudgeInput) -> int:
            return int(judge_input.criteria[0].rsplit(" ", 1)[1])

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal later_peak_concurrent, later_current_concurrent
            first_idx = _first_idx(judge_input)
            attempts[first_idx] = attempts.get(first_idx, 0) + 1
            if first_idx >= 6:
                with lock:
                    later_current_concurrent += 1
                    later_peak_concurrent = max(later_peak_concurrent, later_current_concurrent)
            threading.Event().wait(0.02)
            if first_idx >= 6:
                with lock:
                    later_current_concurrent -= 1
            if first_idx == 2 and attempts[first_idx] == 1:
                return [Verdict(met=None, reasoning="429 rate limit exceeded")], LLMUsage()
            return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()

        mock_run_judge.side_effect = _side_effect

        results, _usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert all(result.met is True for result in results)
        assert attempts[2] == 2
        assert later_peak_concurrent <= 2
        mock_sleep.assert_called_with(1)

    @patch("gandalf.orchestrator.run_judge")
    def test_splits_exceeds_rubric_size(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """batch_splits=5 with 3 criteria → 3 chunks of 1 each."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=5,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(3)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            verdicts = [Verdict(met=True, reasoning=f"ok {i}") for i in range(len(judge_input.criteria))]
            return verdicts, LLMUsage(cost_usd=0.05)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert len(results) == 3
        assert mock_run_judge.call_count == 3
        assert usage.cost_usd == pytest.approx(0.15)

    @patch("gandalf.orchestrator.run_judge")
    def test_trace_file_naming(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Each split gets a unique trace path."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        trace_paths: list[str] = []

        def _side_effect(
            judge_input: BatchJudgeInput, *, trace_path: str, **_kw: Any
        ) -> tuple[list[Verdict], LLMUsage]:
            trace_paths.append(trace_path)
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage()

        mock_run_judge.side_effect = _side_effect

        run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert len(trace_paths) == 2
        assert trace_paths[0] != trace_paths[1]
        # Order may vary due to parallel execution
        trace_basenames = sorted(os.path.basename(p) for p in trace_paths)
        assert trace_basenames[0] == "judge_trace_batch_split0.txt"
        assert trace_basenames[1] == "judge_trace_batch_split1.txt"

    @patch("gandalf.orchestrator.run_judge")
    def test_errored_criteria_in_split(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Errors in one split are properly reflected in merged results."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            # Identify chunk by criteria text — second chunk has "criterion 2"
            is_second_chunk = any("criterion 2" in c for c in judge_input.criteria)
            if not is_second_chunk:
                # First split: both pass
                return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage(cost_usd=0.1)
            # Second split: first criterion errors, second passes
            return [
                Verdict(met=None, reasoning="timeout"),
                Verdict(met=True, reasoning="ok"),
            ], LLMUsage(cost_usd=0.1)

        mock_run_judge.side_effect = _side_effect

        results, _ = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert results[0].met is True
        assert results[1].met is True
        assert results[2].met is None  # errored in second split
        assert results[3].met is True

    @patch("gandalf.orchestrator.run_judge")
    def test_timeout_per_split(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """Each split's timeout is based on its chunk size, not total rubric."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
            judge_timeout=100,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)  # 2 per split

        timeouts: list[int] = []

        def _side_effect(judge_input: BatchJudgeInput, *, timeout: int, **_kw: Any) -> tuple[list[Verdict], LLMUsage]:
            timeouts.append(timeout)
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage()

        mock_run_judge.side_effect = _side_effect

        run_batch_concurrent(config, rubric, "done", "test", "", None)

        # Each split has 2 criteria → timeout = 100 * 2 = 200
        assert all(t == 200 for t in timeouts)

    @patch("gandalf.orchestrator.run_judge")
    def test_batch_timeout_cap_per_split(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """batch_timeout caps each split's timeout independently."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
            judge_timeout=100,
            batch_timeout=150,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        timeouts: list[int] = []

        def _side_effect(judge_input: BatchJudgeInput, *, timeout: int, **_kw: Any) -> tuple[list[Verdict], LLMUsage]:
            timeouts.append(timeout)
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage()

        mock_run_judge.side_effect = _side_effect

        run_batch_concurrent(config, rubric, "done", "test", "", None)

        # 2 criteria * 100s = 200, capped to 150
        assert all(t == 150 for t in timeouts)

    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    @patch("gandalf.orchestrator.run_judge")
    def test_main_dispatches_batch_concurrent(
        self,
        mock_run_judge: Any,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        tmp_path: pathlib.Path,
    ) -> None:
        """main() dispatches to run_batch_concurrent when batch_splits is set."""
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="batch",
            batch_splits=2,
        )
        mock_rubric.return_value = self._make_rubric(4)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage(cost_usd=0.1, prompt_tokens=100, completion_tokens=50)

        mock_run_judge.side_effect = _side_effect

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        info = json.loads((tmp_path / "output" / "info.json").read_text())
        assert len(info["criterion_results"]) == 4
        assert all(r["met"] is True for r in info["criterion_results"])

        reward = json.loads((tmp_path / "output" / "reward.json").read_text())
        assert reward["reward"] == 1.0

    @patch("gandalf.orchestrator.run_judge")
    def test_split_future_raises_exception(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """When run_judge raises an unhandled exception, all criteria fail gracefully."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            # Second chunk has "criterion 2" — that one raises
            is_second_chunk = any("criterion 2" in c for c in judge_input.criteria)
            if not is_second_chunk:
                return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage(cost_usd=0.1)
            msg = "unexpected internal error"
            raise RuntimeError(msg)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        # All criteria should be marked as errored (not just the failed split)
        assert all(r.met is None for r in results)
        assert "Batch split failed" in results[0].reasoning
        # Usage must be reset to stay consistent with all-error results
        assert usage == LLMUsage()

    @patch("gandalf.orchestrator.run_judge")
    def test_split_returns_fewer_verdicts(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """When a split returns fewer verdicts than criteria, missing ones get met=None."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            # Identify chunk by criteria text — second chunk has "criterion 2"
            is_second_chunk = any("criterion 2" in c for c in judge_input.criteria)
            if not is_second_chunk:
                # First split: returns both verdicts
                return [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria], LLMUsage()
            # Second split: only returns 1 verdict for 2 criteria
            return [Verdict(met=True, reasoning="ok")], LLMUsage()

        mock_run_judge.side_effect = _side_effect

        results, _ = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert results[0].met is True  # split 0, verdict present
        assert results[1].met is True  # split 0, verdict present
        assert results[2].met is True  # split 1, position 0 — verdict present
        assert results[3].met is None  # split 1, position 1 — no verdict, defaults to met=None

    @patch("gandalf.orchestrator.run_judge")
    def test_batch_splits_independent_of_max_concurrency(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """batch_splits=4 with max_concurrency=2 creates 4 chunks but only 2 run at a time."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=4,
            max_concurrency=2,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(8)  # 8 criteria / 4 splits = 2 per chunk

        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)  # small delay to overlap threads
            with lock:
                current_concurrent -= 1
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage(cost_usd=0.1)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_batch_concurrent(config, rubric, "done", "test", "", None)

        # 4 splits created (one per chunk)
        assert mock_run_judge.call_count == 4
        # Each chunk should have 2 criteria
        for call in mock_run_judge.call_args_list:
            judge_input = call[0][0]
            assert len(judge_input.criteria) == 2
        # All 8 results present and in order
        assert len(results) == 8
        for i, r in enumerate(results):
            assert r.criterion == f"criterion {i}"
        # Peak concurrency should be capped at 2 (not 4)
        assert peak_concurrent <= 2
        # Total usage from 4 splits
        assert usage.cost_usd == pytest.approx(0.4)

    @patch("gandalf.orchestrator.run_judge")
    def test_fixed_mode_without_max_concurrency_runs_all_splits(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """With discovery disabled, max_concurrency=None runs all explicit splits in parallel."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=3,
            rate_limit_discovery=False,
            # max_concurrency deliberately omitted (None)
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(6)

        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)
            with lock:
                current_concurrent -= 1
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage()

        mock_run_judge.side_effect = _side_effect

        results, _ = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert mock_run_judge.call_count == 3
        assert len(results) == 6
        # All 3 should have been able to run concurrently
        assert peak_concurrent == 3

    @patch("gandalf.orchestrator.run_batch_concurrent")
    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_batch_without_splits_and_auto_parallel_false_dispatches_run_batch(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_batch: Any,
        mock_run_batch_concurrent: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        """mode='batch' can still opt out of automatic splitting."""
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="batch",
            max_concurrency=4,  # should NOT trigger splitting without batch_splits
            auto_parallel=False,
        )
        mock_rubric.return_value = self._make_rubric(4)
        mock_run_batch.return_value = (
            [CriterionResult(criterion=f"criterion {i}", weight=1.0, met=True, reasoning="ok") for i in range(4)],
            LLMUsage(cost_usd=0.1),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        mock_run_batch.assert_called_once()
        mock_run_batch_concurrent.assert_not_called()

    @patch("gandalf.orchestrator.run_batch_concurrent")
    @patch("gandalf.orchestrator.run_batch")
    @patch("gandalf.orchestrator.run_individual")
    @patch("gandalf.orchestrator.resolve_instructions", return_value="test")
    @patch("gandalf.orchestrator.resolve_judge_guidance", return_value="")
    @patch("gandalf.orchestrator.resolve_judge_prompt", return_value=None)
    @patch("gandalf.orchestrator.load_trajectory_final_output", return_value="done")
    @patch("gandalf.orchestrator.load_rubric")
    @patch("gandalf.orchestrator.load_config")
    def test_main_individual_mode_dispatches_run_individual(
        self,
        mock_config: Any,
        mock_rubric: Any,
        mock_trajectory: Any,  # noqa: ARG002
        mock_prompt: Any,  # noqa: ARG002
        mock_guidance: Any,  # noqa: ARG002
        mock_instructions: Any,  # noqa: ARG002
        mock_run_individual: Any,
        mock_run_batch: Any,
        mock_run_batch_concurrent: Any,
        tmp_path: pathlib.Path,
    ) -> None:
        """mode='individual' with max_concurrency dispatches to run_individual."""
        output_dir = str(tmp_path / "output")
        os.makedirs(output_dir, exist_ok=True)

        mock_config.return_value = GraderConfig(
            instructions="test",
            rubric_path="/rubric.json",
            workdir=str(tmp_path),
            trajectory_path="/logs/trajectory.json",
            sandbox_user="sandbox",
            output_dir=output_dir,
            mode="individual",
            max_concurrency=3,
        )
        mock_rubric.return_value = self._make_rubric(4)
        mock_run_individual.return_value = (
            [CriterionResult(criterion=f"criterion {i}", weight=1.0, met=True, reasoning="ok") for i in range(4)],
            LLMUsage(cost_usd=0.1),
        )

        with patch("sys.argv", ["prog", "--config", "dummy.toml"]):
            main()

        mock_run_individual.assert_called_once()
        mock_run_batch.assert_not_called()
        mock_run_batch_concurrent.assert_not_called()

    @patch("gandalf.orchestrator.run_judge")
    def test_max_concurrency_capped_to_chunk_count(self, mock_run_judge: Any, tmp_path: pathlib.Path) -> None:
        """batch_splits=2, max_concurrency=10 — thread pool capped to 2 (the actual chunk count)."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="batch",
            batch_splits=2,
            max_concurrency=10,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(4)

        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(judge_input: BatchJudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)
            with lock:
                current_concurrent -= 1
            verdicts = [Verdict(met=True, reasoning="ok") for _ in judge_input.criteria]
            return verdicts, LLMUsage(cost_usd=0.1)

        mock_run_judge.side_effect = _side_effect

        results, _ = run_batch_concurrent(config, rubric, "done", "test", "", None)

        assert mock_run_judge.call_count == 2
        assert len(results) == 4
        # Peak concurrency must be 2 (chunks), not 10 (max_concurrency)
        assert peak_concurrent <= 2

    @patch("gandalf.orchestrator.run_judge")
    def test_individual_mode_parallelizes_with_max_concurrency(
        self, mock_run_judge: Any, tmp_path: pathlib.Path
    ) -> None:
        """run_individual with max_concurrency=3 actually evaluates criteria concurrently."""
        config = make_config(
            workdir=str(tmp_path),
            output_dir=str(tmp_path / "output"),
            mode="individual",
            max_concurrency=3,
        )
        os.makedirs(config.output_dir, exist_ok=True)
        rubric = self._make_rubric(6)

        peak_concurrent = 0
        current_concurrent = 0
        lock = threading.Lock()

        def _side_effect(_judge_input: JudgeInput, **_kwargs: Any) -> tuple[list[Verdict], LLMUsage]:
            nonlocal peak_concurrent, current_concurrent
            with lock:
                current_concurrent += 1
                peak_concurrent = max(peak_concurrent, current_concurrent)
            time.sleep(0.05)
            with lock:
                current_concurrent -= 1
            return [Verdict(met=True, reasoning="ok")], LLMUsage(cost_usd=0.01)

        mock_run_judge.side_effect = _side_effect

        results, usage = run_individual(config, rubric, "done", "test", "", None)

        # All 6 criteria evaluated individually
        assert mock_run_judge.call_count == 6
        assert len(results) == 6
        for i, r in enumerate(results):
            assert r.criterion == f"criterion {i}"
            assert r.met is True
        # Must have actually parallelized — peak > 1
        assert peak_concurrent > 1
        assert peak_concurrent <= 3
        assert usage.cost_usd == pytest.approx(0.06)
